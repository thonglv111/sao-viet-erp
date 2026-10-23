"""Hoa hồng kinh doanh — tính theo TIỀN THU VỀ, % theo biểu bậc số ngày thu (chủ chốt 23/10/2026,
spec 2026-10-23-hoa-hong-theo-tien-thu-design.md; đảo mốc "hoá đơn" của 21/08/2026).

Khoá các luật đổi SỐ TIỀN:
  · MỐC = dòng phiếu thu `received` có `receipt_date` trong kỳ;
  · GỐC = tiền thu × tỷ lệ TRƯỚC VAT − phần CK + kê của hoá đơn theo tỷ lệ tiền thu;
  · % = `orders.hh_rieng_pct` nếu có, không thì biểu đã CHỤP vào đơn tra theo số ngày thu.
Số đối chiếu lấy từ sổ lương T08/2026 (sheet DOANH THU T8, Lương KD CTV).
"""
from __future__ import annotations

from datetime import date

from app.db import SessionLocal
from app.models.accounting import (
    PAYMENT_RECEIPT_CANCELLED,
    PAYMENT_RECEIPT_RECEIVED,
    RECEIPT_SOURCE_SALES_INVOICE,
    SALES_INVOICE_CANCELLED,
    SALES_INVOICE_ISSUED,
    PaymentReceipt,
    PaymentReceiptLine,
    SalesInvoice,
)
from app.models.customer import Customer
from app.models.employee import STATUS_ACTIVE, Employee
from app.models.order import STATUS_ORDERED, Order, OrderLine
from app.models.payroll import EmployeeSalary
from app.models.user import User
from app.security import hash_password
from app.services.hoa_hong_service import HoaHongService

KY_TU, KY_DEN = date(2026, 8, 1), date(2026, 8, 31)


def _sales(ten: str, *, pct: float = 0.05, bac: list | None = None) -> tuple[int, int]:
    """NV kinh doanh có tài khoản + mức lương khai `commission_pct`. Trả `(employee_id, user_id)`."""
    db = SessionLocal()
    try:
        u = User(username=f"sale-{ten}", name=ten, password_hash=hash_password("x"))
        db.add(u)
        db.flush()
        e = Employee(code=f"NVKD{abs(hash(ten)) % 9000 + 1000}", full_name=ten,
                     hire_date=date(2020, 1, 1), user_id=u.id, status=STATUS_ACTIVE)
        db.add(e)
        db.flush()
        db.add(EmployeeSalary(employee_id=e.id, effective_from=date(2020, 1, 1),
                              base_amount=10_000_000, commission_pct=pct, bac_hoa_hong=bac))
        db.commit()
        return e.id, u.id
    finally:
        db.close()


def _don(ma: str, user_id: int, *, truoc_vat: int, vat_pct: int = 8,
         pct: float | None = 0.05, bac: list | None = None, chiet_khau: int = 0, ke: int = 0) -> int:
    """Đơn ĐÃ CHỐT của người sales đó, biểu đã chụp = một bậc vô hạn `pct` (hoặc `bac`).
    `pct=None` và không `bac` ⇒ đơn không có hoa hồng."""
    db = SessionLocal()
    try:
        kh = Customer(code=f"KH-{ma}", name=f"Khach {ma}")
        db.add(kh)
        db.flush()
        bieu = bac if bac is not None else ([{"den_ngay": None, "pct": pct}] if pct else None)
        o = Order(order_no=ma, customer_id=kh.id, status=STATUS_ORDERED,
                  sale_user_id=user_id, bac_hoa_hong=bieu)
        o.lines.append(OrderLine(description="Hop giay", qty=1, don_vi_tinh="cai",
                                 line_total=truoc_vat, vat_pct_estimate=vat_pct,
                                 chiet_khau=chiet_khau or None, ke_gia=ke or None))
        db.add(o)
        db.commit()
        return o.id
    finally:
        db.close()


def _hoa_don_chi(order_id: int, so_tien_co_vat: int, *, ngay: date = date(2026, 8, 15),
                 chiet_khau: int = 0, ke: int = 0) -> int:
    """Hoá đơn issued, CHƯA thu đồng nào. Trả id."""
    db = SessionLocal()
    try:
        o = db.get(Order, order_id)
        inv = SalesInvoice(
            order_id=order_id, customer_id=o.customer_id,
            invoice_symbol="1C26TAA", invoice_number=f"{order_id}{ngay:%m%d}{so_tien_co_vat % 997}",
            invoice_date=ngay, amount_vnd=so_tien_co_vat, chiet_khau=chiet_khau, ke_gia=ke,
            customer_name_snapshot="Khach test", status=SALES_INVOICE_ISSUED,
        )
        db.add(inv)
        db.commit()
        return inv.id
    finally:
        db.close()


_SO_PHIEU = [0]


def _thu(dong: list[tuple[int, int | None, int]], *, ngay: date = date(2026, 8, 15),
         huy: bool = False) -> int:
    """MỘT phiếu thu, mỗi dòng `(order_id, sales_invoice_id | None = cọc, số tiền)`. Trả id phiếu."""
    db = SessionLocal()
    try:
        _SO_PHIEU[0] += 1
        o = db.get(Order, dong[0][0])
        tong = sum(t for _o, _i, t in dong)
        r = PaymentReceipt(
            code=f"PT-HH-{_SO_PHIEU[0]:04d}", source_type=RECEIPT_SOURCE_SALES_INVOICE,
            order_id=o.id, customer_id=o.customer_id, payer_name="Khach test", receipt_method="cash",
            status=PAYMENT_RECEIPT_CANCELLED if huy else PAYMENT_RECEIPT_RECEIVED,
            receipt_date=ngay, amount=tong, amount_vnd=tong, content="Thu tien hang",
        )
        for oid, iid, tien in dong:
            r.lines.append(PaymentReceiptLine(order_id=oid, sales_invoice_id=iid, amount_vnd=tien))
        db.add(r)
        db.commit()
        return r.id
    finally:
        db.close()


def _hoa_don(order_id: int, so_tien_co_vat: int, *, ngay: date = date(2026, 8, 15),
             huy: bool = False) -> int:
    """Hoá đơn + THU ĐỦ ngay hôm đó (0 ngày). `huy` ⇒ phiếu thu bị huỷ. Trả id hoá đơn."""
    iid = _hoa_don_chi(order_id, so_tien_co_vat, ngay=ngay)
    _thu([(order_id, iid, so_tien_co_vat)], ngay=ngay, huy=huy)
    return iid


def _tinh(employee_id: int, *, tu=KY_TU, den=KY_DEN) -> float:
    db = SessionLocal()
    try:
        return HoaHongService(db).hoa_hong_ky(employee_id, tu_ngay=tu, den_ngay=den)
    finally:
        db.close()


# =============================================================================================
# Công thức
# =============================================================================================
def test_THU_TIEN_la_co_hoa_hong(client):
    """⭐ Đơn 100tr trước VAT + VAT 8% ⇒ thu 108tr; hoa hồng 5% trên phần TRƯỚC VAT = 5tr."""
    emp, uid = _sales("ra cong no")
    oid = _don("DH-HH-01", uid, truoc_vat=100_000_000, vat_pct=8)
    _hoa_don(oid, 108_000_000)
    assert _tinh(emp) == 5_000_000


def test_GOC_TINH_la_TRUOC_VAT_khong_phai_tien_hoa_don(client):
    """⭐ VAT là tiền thu hộ nhà nước — trả hoa hồng trên đó là trả trên tiền không phải của mình.

    Nếu lấy thẳng tiền hoá đơn thì ra 108tr × 5% = 5.4tr, thừa 400k cho mỗi 100tr doanh thu.
    """
    emp, uid = _sales("truoc vat")
    oid = _don("DH-HH-02", uid, truoc_vat=100_000_000, vat_pct=8)
    _hoa_don(oid, 108_000_000)
    assert _tinh(emp) == 5_000_000
    assert _tinh(emp) != 5_400_000


def test_don_KHONG_VAT_thi_khong_quy_doi_gi(client):
    """Tỷ lệ quy đổi phải là 1.0, không được chia cho 0 hay cắt mất tiền."""
    emp, uid = _sales("khong vat")
    oid = _don("DH-HH-03", uid, truoc_vat=50_000_000, vat_pct=0)
    _hoa_don(oid, 50_000_000)
    assert _tinh(emp) == 2_500_000


def test_HOA_DON_TUNG_PHAN_chi_an_hoa_hong_phan_do(client):
    """⭐ Một đơn xuất nhiều hoá đơn ⇒ mỗi hoá đơn chỉ mở phần hoa hồng của nó.

    Đây là lý do quy đổi theo TỶ LỆ chứ không trừ thẳng VAT: hoá đơn chỉ là một mẩu của đơn.
    """
    emp, uid = _sales("tung phan")
    oid = _don("DH-HH-04", uid, truoc_vat=100_000_000, vat_pct=8)
    _hoa_don(oid, 54_000_000, ngay=date(2026, 8, 10))    # nửa đơn
    assert _tinh(emp) == 2_500_000
    _hoa_don(oid, 54_000_000, ngay=date(2026, 8, 20))    # nửa còn lại
    assert _tinh(emp) == 5_000_000


# =============================================================================================
# Các trường KHÔNG được tính
# =============================================================================================
def test_hoa_hong_tru_ck_va_toan_bo_ke(client):
    """CK và kê của hoá đơn chi lại khách (20/10/2026) ⇒ gốc tính trừ CK và TOÀN BỘ kê (cả phần thuế
    TNDN trên kê). HĐ 16,5tr gồm VAT 10%, CK 1tr, kê 5tr ⇒ (15tr − 6tr) × 5% = 450k."""
    emp, uid = _sales("tru ck ke")
    oid = _don("DH-HH-CK", uid, truoc_vat=15_000_000, vat_pct=10)
    _hoa_don(oid, 16_500_000)
    db = SessionLocal()
    try:
        inv = db.query(SalesInvoice).filter(SalesInvoice.order_id == oid).one()
        inv.chiet_khau, inv.ke_gia, inv.ke_thue_tndn = 1_000_000, 5_000_000, 1_000_000
        db.commit()
    finally:
        db.close()
    assert _tinh(emp) == 450_000


def test_hoa_don_HUY_khong_tinh(client):
    emp, uid = _sales("huy")
    oid = _don("DH-HH-05", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000, huy=True)
    assert _tinh(emp) == 0


def test_hoa_don_NGOAI_KY_khong_tinh(client):
    """Hoá đơn tháng 7 không rơi vào kỳ tháng 8, và ngược lại."""
    emp, uid = _sales("ngoai ky")
    oid = _don("DH-HH-06", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000, ngay=date(2026, 7, 31))
    assert _tinh(emp) == 0
    assert _tinh(emp, tu=date(2026, 7, 1), den=date(2026, 7, 31)) == 5_000_000


def test_don_cua_NGUOI_KHAC_khong_tinh(client):
    """⭐ Hoa hồng bám `orders.sale_user_id`. Sai chỗ này là trả tiền cho nhầm người."""
    a, uid_a = _sales("nguoi A")
    b, _uid_b = _sales("nguoi B")
    oid = _don("DH-HH-07", uid_a, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    assert _tinh(a) == 5_000_000
    assert _tinh(b) == 0


def test_don_KHONG_KHAI_hoa_hong_thi_bang_0(client):
    """Đơn không chụp biểu ⇒ không có hoa hồng, dù người sales có biểu riêng."""
    emp, uid = _sales("don khong khai")
    oid = _don("DH-HH-08", uid, truoc_vat=100_000_000, pct=None)
    _hoa_don(oid, 108_000_000)
    assert _tinh(emp) == 0


def test_pct_AM_khong_duoc_TRU_vao_hoa_hong_don_khac(client):
    """% âm (gõ nhầm dấu) chỉ được BỎ QUA đơn đó, không được ăn lẹm sang đơn lành.

    Không có hàng rào `pct <= 0` thì đơn âm trừ thẳng vào tổng: người ta làm 2 đơn mà nhận ít hơn
    làm 1 đơn, và trên phiếu lương không có gì để lần ra vì sao.
    """
    emp, uid = _sales("pct am")
    lanh = _don("DH-HH-11", uid, truoc_vat=100_000_000)
    _hoa_don(lanh, 108_000_000, ngay=date(2026, 8, 11))
    am = _don("DH-HH-12", uid, truoc_vat=100_000_000, pct=-0.05)
    _hoa_don(am, 108_000_000, ngay=date(2026, 8, 12))
    assert _tinh(emp) == 5_000_000, "đơn % âm đã ăn lẹm vào hoa hồng của đơn lành"


def test_nhan_vien_CHUA_CO_TAI_KHOAN_thi_bang_0(client):
    """Đơn ghi USER, bảng lương chạy theo EMPLOYEE — không có tài khoản thì không thể là sales."""
    db = SessionLocal()
    try:
        e = Employee(code="NVKD-NOUSER", full_name="Khong co tai khoan",
                     hire_date=date(2020, 1, 1))
        db.add(e)
        db.commit()
        eid = e.id
    finally:
        db.close()
    assert _tinh(eid) == 0


def test_don_khong_co_dong_hang_khong_lam_vo(client):
    """Đơn chưa có dòng nào ⇒ không biết phần trước VAT ⇒ gốc = 0 (plan Review Focus 3), không chia
    cho 0. Dữ liệu hỏng thì không trả tiền đoán."""
    db = SessionLocal()
    try:
        u = User(username="sale-rong", name="Sale rong", password_hash=hash_password("x"))
        db.add(u)
        db.flush()
        e = Employee(code="NVKD-RONG", full_name="Sale rong", hire_date=date(2020, 1, 1),
                     user_id=u.id)
        db.add(e)
        db.flush()
        kh = Customer(code="KH-RONG", name="Khach rong")
        db.add(kh)
        db.flush()
        o = Order(order_no="DH-HH-09", customer_id=kh.id, status=STATUS_ORDERED,
                  sale_user_id=u.id, bac_hoa_hong=[{"den_ngay": None, "pct": 0.05}])
        db.add(o)
        db.commit()
        eid, oid = e.id, o.id
    finally:
        db.close()
    _hoa_don(oid, 10_000_000)
    assert _tinh(eid) == 0


# =============================================================================================
# Theo tiền thu về + biểu bậc số ngày (23/10/2026) — số đối chiếu sổ T08
# =============================================================================================
LANH_DAO = [{"den_ngay": None, "pct": 0.005}]
CTV = [{"den_ngay": 45, "pct": 0.03}, {"den_ngay": 60, "pct": 0.02},
       {"den_ngay": 75, "pct": 0.01}, {"den_ngay": None, "pct": 0}]


def test_K_DUNG_GOODY_124_ngay_ra_259445(client):
    """Sổ T08 dòng 249: HĐ 25/04 thu 27/08, 56.040.120 ÷ 1,08 = 51.889.000 × 0,5% = 259.445."""
    emp, uid = _sales("k dung")
    oid = _don("DH-GOODY", uid, truoc_vat=51_889_000, vat_pct=8, bac=LANH_DAO)
    iid = _hoa_don_chi(oid, 56_040_120, ngay=date(2026, 4, 25))
    _thu([(oid, iid, 56_040_120)], ngay=date(2026, 8, 27))
    assert _tinh(emp) == 259_445


def test_CTV_PINACO_30_ngay_tru_CK_ra_1423464(client):
    """Sổ T08 (Lương KD CTV): HĐ 23/07 thu 22/08, 53.620.704 ÷ 1,08 − CK 2.200.000 = 47.448.800 × 3%."""
    emp, uid = _sales("n son")
    oid = _don("DH-PINACO", uid, truoc_vat=49_648_800, vat_pct=8, bac=CTV)
    iid = _hoa_don_chi(oid, 53_620_704, ngay=date(2026, 7, 23), chiet_khau=2_200_000)
    _thu([(oid, iid, 53_620_704)], ngay=date(2026, 8, 22))
    assert _tinh(emp) == 1_423_464


def test_GOC_AM_cua_mot_lan_thu_tru_vao_tong_nhu_so(client):
    """Sổ T08 dòng 323 (S SƠN Đ, NAM HOA): thu 93.085.200 mà CK 321.781.872 ⇒ doanh thu thực
    −235.591.872, sổ cộng thẳng số âm vào tổng của Sơn (4.677.451.077 × 0,5% = 23.387.255). Kẹp gốc
    từng lần thu về 0 thì Sơn dư 1.177.959 so với sổ. Tổng của người vẫn không xuống dưới 0."""
    emp, uid = _sales("s son")
    lon = _don("DH-LON", uid, truoc_vat=1_000_000_000, vat_pct=8, bac=LANH_DAO)
    _hoa_don(lon, 1_080_000_000, ngay=date(2026, 8, 10))
    am = _don("DH-NAM-HOA", uid, truoc_vat=86_190_000, vat_pct=8, bac=LANH_DAO)
    iid = _hoa_don_chi(am, 93_085_200, ngay=date(2026, 8, 1), chiet_khau=321_781_872)
    _thu([(am, iid, 93_085_200)], ngay=date(2026, 8, 20))
    assert _tinh(emp) == 3_822_041         # 5.000.000 − 235.591.872 × 0,5%


def test_moc_la_THANG_THU_khong_phai_thang_hoa_don(client):
    emp, uid = _sales("thang thu")
    oid = _don("DH-THANG-THU", uid, truoc_vat=100_000_000)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2026, 7, 20))
    _thu([(oid, iid, 108_000_000)], ngay=date(2026, 8, 5))
    assert _tinh(emp, tu=date(2026, 7, 1), den=date(2026, 7, 31)) == 0
    assert _tinh(emp) == 5_000_000


def _ctv_thu_sau(ngay_hd: date, ma: str) -> int:
    emp, uid = _sales(f"ctv {ma}")
    oid = _don(ma, uid, truoc_vat=100_000_000, bac=CTV)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=ngay_hd)
    _thu([(oid, iid, 108_000_000)], ngay=date(2026, 8, 20))
    return _tinh(emp)


def test_bac_45_ngay_ra_3pct_46_ngay_ra_2pct(client):
    assert _ctv_thu_sau(date(2026, 7, 6), "DH-B45") == 3_000_000     # 45 ngày
    assert _ctv_thu_sau(date(2026, 7, 5), "DH-B46") == 2_000_000     # 46 ngày


def test_76_ngay_ra_0(client):
    assert _ctv_thu_sau(date(2026, 6, 5), "DH-B76") == 0


def test_lanh_dao_khoan_383_ngay_van_tinh(client):
    emp, uid = _sales("s son l")
    oid = _don("DH-383", uid, truoc_vat=100_000_000, bac=LANH_DAO)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2025, 8, 2))
    _thu([(oid, iid, 108_000_000)], ngay=date(2026, 8, 20))
    assert _tinh(emp) == 500_000


def test_thu_TRUOC_hoa_don_tinh_0_ngay(client):
    """Sổ T08 dòng 395: thu 15/08, HĐ 27/08 (−12 ngày) ⇒ bậc đầu, không văng lỗi."""
    emp, uid = _sales("thu truoc")
    oid = _don("DH-TRUOC", uid, truoc_vat=100_000_000, bac=CTV)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2026, 8, 27))
    _thu([(oid, iid, 108_000_000)], ngay=date(2026, 8, 15))
    assert _tinh(emp) == 3_000_000


def test_thu_TUNG_PHAN_hai_ky_chia_CK_ke_theo_ty_le(client):
    """HĐ 108tr có CK 5tr, thu 54tr tháng 8 + 54tr tháng 9 ⇒ mỗi kỳ gốc 50tr − 2,5tr = 47,5tr."""
    emp, uid = _sales("tung phan ck")
    oid = _don("DH-TP-CK", uid, truoc_vat=100_000_000)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2026, 8, 1), chiet_khau=5_000_000)
    _thu([(oid, iid, 54_000_000)], ngay=date(2026, 8, 10))
    _thu([(oid, iid, 54_000_000)], ngay=date(2026, 9, 10))
    assert _tinh(emp) == 2_375_000
    assert _tinh(emp, tu=date(2026, 9, 1), den=date(2026, 9, 30)) == 2_375_000


def test_coc_tru_CK_ke_theo_ty_le_don_va_0_ngay(client):
    """Cọc 54tr trên đơn 100tr trước VAT (VAT 8%), dòng đơn có CK 4tr + kê 6tr:
    gốc = 54tr × (100tr − 10tr) / 108tr = 45tr; biểu CTV, 0 ngày ⇒ 3%."""
    emp, uid = _sales("coc")
    oid = _don("DH-COC", uid, truoc_vat=100_000_000, bac=CTV, chiet_khau=4_000_000, ke=6_000_000)
    _thu([(oid, None, 54_000_000)], ngay=date(2026, 8, 3))
    assert _tinh(emp) == 1_350_000


def test_phieu_thu_HUY_khong_tinh(client):
    emp, uid = _sales("phieu huy")
    oid = _don("DH-PT-HUY", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000, huy=True)
    assert _tinh(emp) == 0


def _dat_rieng(oid: int, pct: float | None) -> None:
    db = SessionLocal()
    try:
        db.get(Order, oid).hh_rieng_pct = pct
        db.commit()
    finally:
        db.close()


def test_hh_rieng_thang_bieu(client):
    """Đơn Wipro 1% thắng biểu 3%."""
    emp, uid = _sales("wipro")
    oid = _don("DH-WIPRO", uid, truoc_vat=100_000_000, bac=CTV)
    _hoa_don(oid, 108_000_000)
    _dat_rieng(oid, 0.01)
    assert _tinh(emp) == 1_000_000


def test_hh_rieng_bang_0_la_loai_tay(client):
    """`hh_rieng_pct = 0` (khác NULL) ⇒ đơn ra 0 dù biểu có % — loại tay khoản KARTA."""
    emp, uid = _sales("karta")
    oid = _don("DH-KARTA", uid, truoc_vat=100_000_000, bac=LANH_DAO)
    _hoa_don(oid, 108_000_000)
    _dat_rieng(oid, 0)
    assert _tinh(emp) == 0


def test_doi_bieu_sau_khi_chot_don_cu_giu_bieu_da_chup(client):
    """Đơn đã chốt giữ biểu cũ — tiền đã hứa thì không sửa ngược."""
    emp, uid = _sales("doi bieu", bac=[{"den_ngay": None, "pct": 0.05}])
    oid = _don("DH-DOI-BIEU", uid, truoc_vat=100_000_000, pct=0.05)
    _hoa_don(oid, 108_000_000)
    db = SessionLocal()
    try:
        db.add(EmployeeSalary(employee_id=emp, effective_from=date(2026, 8, 1), base_amount=10_000_000,
                              bac_hoa_hong=[{"den_ngay": None, "pct": 0.20}]))
        db.commit()
    finally:
        db.close()
    assert _tinh(emp) == 5_000_000, "đổi biểu của người làm đổi luôn hoa hồng đơn đã chốt"


def test_mot_phieu_nhieu_don_khac_nguoi(client):
    """Một phiếu thu trả cho hai đơn của hai người ⇒ mỗi dòng về đúng người của đơn đó."""
    a, ua = _sales("phieu A")
    b, ub = _sales("phieu B")
    oa = _don("DH-PA", ua, truoc_vat=100_000_000)
    ob = _don("DH-PB", ub, truoc_vat=50_000_000)
    ia = _hoa_don_chi(oa, 108_000_000)
    ib = _hoa_don_chi(ob, 54_000_000)
    _thu([(oa, ia, 108_000_000), (ob, ib, 54_000_000)])
    assert _tinh(a) == 5_000_000
    assert _tinh(b) == 2_500_000


def test_hoa_don_so_tien_0_khong_chia_0(client):
    """Hoá đơn 0 đồng mà vẫn có tiền thu trỏ vào ⇒ dữ liệu hỏng ⇒ gốc = 0 (plan Review Focus 3)."""
    emp, uid = _sales("hd 0")
    oid = _don("DH-HD0", uid, truoc_vat=100_000_000)
    iid = _hoa_don_chi(oid, 0, chiet_khau=1_000_000)
    _thu([(oid, iid, 10_800_000)])
    assert _tinh(emp) == 0


def _huy_hoa_don(iid: int) -> None:
    db = SessionLocal()
    try:
        db.get(SalesInvoice, iid).status = SALES_INVOICE_CANCELLED
        db.commit()
    finally:
        db.close()


def test_thu_tren_hoa_don_DA_HUY_tinh_goc_theo_don_nhu_tien_coc(client):
    """Hoá đơn huỷ khi khách đã trả: phiếu thu giữ nguyên, tiền thành quỹ của đơn (Kế toán 17/10).
    CK/kê của hoá đơn huỷ không còn giá trị ⇒ gốc trừ CK/kê theo tỷ lệ CỦA ĐƠN như tiền cọc; số ngày
    vẫn tính từ ngày hoá đơn huỷ (khách nợ thật từ hôm đó)."""
    emp, uid = _sales("hd huy", bac=CTV)
    oid = _don("DH-HD-HUY", uid, truoc_vat=100_000_000, bac=CTV, chiet_khau=2_000_000)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2026, 7, 5), chiet_khau=50_000_000)
    _thu([(oid, iid, 108_000_000)], ngay=date(2026, 8, 10))        # 36 ngày ⇒ 3%
    _huy_hoa_don(iid)
    # 108tr × (100tr − 2tr) / 108tr = 98tr × 3%
    assert _tinh(emp) == 2_940_000
    db = SessionLocal()
    try:
        k = HoaHongService(db).cac_khoan(KY_TU, KY_DEN)[0]
    finally:
        db.close()
    assert (k["so_ngay"], k["hoa_don_huy"]) == (36, True)


# =============================================================================================
# Chụp ảnh % lúc chốt đơn
# =============================================================================================
def test_CHOT_DON_chup_bieu_cua_sale_vao_don(client):
    """⭐ Chụp lúc chốt, không đọc-sống: đổi biểu tháng sau không được sửa ngược đơn cũ."""
    from app.services.order_service import _bac_hoa_hong_cua_sale

    _emp, uid = _sales("chup anh", bac=[{"den_ngay": None, "pct": 0.07}])
    db = SessionLocal()
    try:
        assert _bac_hoa_hong_cua_sale(db, uid) == [{"den_ngay": None, "pct": 0.07}]
        assert _bac_hoa_hong_cua_sale(db, None) is None
        assert _bac_hoa_hong_cua_sale(db, 999999) is None
    finally:
        db.close()


def test_CHOT_DON_ghi_bieu_len_don(client):
    """Chốt đơn qua API ⇒ `orders.bac_hoa_hong` = biểu của người phụ trách tại ngày chốt."""
    from app.services.order_service import OrderService

    bieu = [{"den_ngay": 45, "pct": 0.03}, {"den_ngay": None, "pct": 0}]
    _emp, uid = _sales("chot ghi", bac=bieu)
    db = SessionLocal()
    try:
        kh = Customer(code="KH-CHOT-BIEU", name="Khach chot")
        db.add(kh)
        db.flush()
        o = Order(order_no="DH-CHOT-BIEU", customer_id=kh.id, status="draft", sale_user_id=uid)
        db.add(o)
        db.commit()
        o.bac_hoa_hong = None
        from app.services import order_service as osv
        osv._chup_hoa_hong(o, db)
        assert o.bac_hoa_hong == bieu
    finally:
        db.close()


def test_chup_bieu_KHONG_duoc_lay_muc_CUA_TUONG_LAI(client):
    """⭐ Chốt đơn HÔM NAY phải ăn biểu ĐANG hiệu lực hôm nay, không phải biểu tháng sau."""
    from app.services.order_service import _bac_hoa_hong_cua_sale

    emp, uid = _sales("muc tuong lai", bac=[{"den_ngay": None, "pct": 0.05}])
    db = SessionLocal()
    try:
        db.add(EmployeeSalary(employee_id=emp, effective_from=date(2099, 12, 1), base_amount=10_000_000,
                              bac_hoa_hong=[{"den_ngay": None, "pct": 0.20}]))
        db.commit()
        assert _bac_hoa_hong_cua_sale(db, uid) == [{"den_ngay": None, "pct": 0.05}], (
            "chụp nhầm biểu của tương lai — đơn chốt hôm nay ăn % chưa tới hạn")
    finally:
        db.close()


def test_chua_khai_bieu_thi_chup_None(client):
    from app.models.user import User as _U
    from app.services.order_service import _bac_hoa_hong_cua_sale

    db = SessionLocal()
    try:
        u = _U(username="sale-chua-khai", name="Chua khai", password_hash=hash_password("x"))
        db.add(u)
        db.flush()
        db.add(Employee(code="NVKD-CHUAKHAI", full_name="Chua khai", hire_date=date(2020, 1, 1),
                        user_id=u.id, status=STATUS_ACTIVE))
        db.commit()
        assert _bac_hoa_hong_cua_sale(db, u.id) is None
    finally:
        db.close()


# =============================================================================================
# Nối vào BẢNG LƯƠNG — đây mới là chỗ tiền tới tay người ta
# =============================================================================================
ADMIN = {"username": "admin", "password": "admin123"}


def _h(client) -> dict[str, str]:
    tok = client.post("/api/auth/login", json=ADMIN).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _tinh_luong(client, eid: int, *, year=2026, month=8) -> dict:
    """Bấm "Tính lại" cả kỳ rồi lấy dòng của NV này."""
    r = client.post("/api/luong/generate", json={"year": year, "month": month}, headers=_h(client))
    assert r.status_code in (200, 201), r.text
    return next(l for l in r.json()["lines"] if l["employee_id"] == eid)


def test_hoa_hong_LEN_PHIEU_LUONG_va_cong_vao_gross(client):
    """⭐ Cả tính năng chỉ có nghĩa ở đây: % khai xong phải RA TIỀN trên phiếu lương.

    Từ 07/09/2026 hoa hồng là CỘT `hoa_hong` trên dòng lương (chủ: "nó là một dạng lương"), không
    còn là dòng khoản danh mục nguồn `auto`.
    """
    emp, uid = _sales("len phieu")
    truoc = _tinh_luong(client, emp)                  # chưa có hoá đơn nào
    assert truoc["hoa_hong"] == 0
    oid = _don("DH-LUONG-01", uid, truoc_vat=100_000_000, vat_pct=8)
    _hoa_don(oid, 108_000_000)
    sau = _tinh_luong(client, emp)
    assert sau["hoa_hong"] == 5_000_000
    assert sau["gross"] == truoc["gross"] + 5_000_000, (
        f"hoa hồng có số nhưng gross đi từ {truoc['gross']} sang {sau['gross']}")
    # Không còn dòng khoản danh mục nào mang hoa hồng.
    assert all(c["code"] != "hoa_hong_kd" for c in sau["components"])


def test_KHONG_de_dong_0_dong_cho_nguoi_khong_lam_kinh_doanh(client):
    """Người không làm kinh doanh: cột hoa hồng = 0 và KHÔNG có dòng khoản rác nào trên phiếu."""
    emp, _uid = _sales("khong kinh doanh")
    line = _tinh_luong(client, emp)
    assert line["hoa_hong"] == 0
    assert all(c["code"] != "hoa_hong_kd" for c in line["components"])


def test_TINH_LAI_khong_cong_doi(client):
    """⭐ Bẫy chết người của khoản hệ tự tính: mỗi lần bấm "Tính lại" lại cộng dồn."""
    emp, uid = _sales("tinh lai")
    oid = _don("DH-LUONG-02", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    l1 = _tinh_luong(client, emp)
    l2 = _tinh_luong(client, emp)
    l3 = _tinh_luong(client, emp)
    assert l3["hoa_hong"] == 5_000_000, "tính lại 3 lần cộng dồn hoa hồng"
    assert l3["gross"] == l1["gross"] == l2["gross"]


def test_HOA_DON_MOI_thi_so_moi_THAY_so_cu(client):
    """Xuất thêm hoá đơn giữa kỳ ⇒ tính lại phải ra số MỚI: không giữ số cũ, cũng không cộng dồn."""
    emp, uid = _sales("hoa don moi")
    oid = _don("DH-LUONG-03", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 54_000_000, ngay=date(2026, 8, 10))
    assert _tinh_luong(client, emp)["hoa_hong"] == 2_500_000
    _hoa_don(oid, 54_000_000, ngay=date(2026, 8, 20))
    assert _tinh_luong(client, emp)["hoa_hong"] == 5_000_000, (
        "phải là 5tr (số mới), không phải 2,5tr hay 7,5tr")


def _dong_bang(client, line_id: int, *, year=2026, month=8) -> dict:
    """Đọc lại dòng từ bảng lương (KHÔNG tính lại) — để soi đường "Sửa 1 ô"."""
    r = client.get("/api/luong/table", params={"year": year, "month": month}, headers=_h(client))
    assert r.status_code == 200, r.text
    return next(l for l in r.json()["lines"] if l["id"] == line_id)


def test_tinh_lai_hoa_hong_KHONG_XOA_thuong_nong_them_tay_va_SUA_1_O_khong_lam_mat_hoa_hong(client):
    """⭐ Hai bẫy trên cùng một dòng:
    1. "Tính lại" phải giữ thưởng nóng HCNS đã thêm tay (nguồn `line`).
    2. Thêm khoản phát sinh đi qua đường "Sửa 1 ô" (`update_line`). Trước 07/09/2026 đường này
       KHÔNG cộng nguồn `auto` ⇒ ngay lúc bấm Thêm, gross tụt mất 5tr hoa hồng mà bảng lương vẫn
       trông bình thường (bản rà 07/09, D1). Nay hoa hồng là cột: gross chỉ được tăng đúng 300k.
    """
    emp, uid = _sales("giu thuong nong")
    oid = _don("DH-LUONG-04", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    line = _tinh_luong(client, emp)
    assert line["hoa_hong"] == 5_000_000
    comps = client.get("/api/luong/components", headers=_h(client)).json()["items"]
    khac = next(c for c in comps if c["code"] == "thu_nhap_khac_ct")
    r = client.post(f"/api/luong/lines/{line['id']}/components",
                    json={"component_id": khac["id"], "amount": 300_000}, headers=_h(client))
    assert r.status_code == 201, r.text
    sua = _dong_bang(client, line["id"])               # số ĐANG ghi sau "Sửa 1 ô", chưa Tính lại
    assert sua["hoa_hong"] == 5_000_000
    assert sua["gross"] == line["gross"] + 300_000, (
        f"Sửa 1 ô làm gross đi từ {line['gross']} sang {sua['gross']} — hoa hồng bốc hơi?")
    sau = _tinh_luong(client, emp)                     # bấm "Tính lại"
    ma = {c["code"]: c["amount"] for c in sau["components"]}
    assert ma.get("thu_nhap_khac_ct") == 300_000, "tính lại đã nuốt mất thưởng nóng"
    assert sau["hoa_hong"] == 5_000_000
    assert sau["gross"] == sua["gross"], "Tính lại và Sửa 1 ô phải ra cùng một gross"


def test_hoa_hong_CHIU_THUE_TNCN(client):
    """Hoa hồng là thu nhập từ tiền lương ⇒ LUÔN chịu thuế (cột riêng, không còn cờ danh mục)."""
    emp, uid = _sales("chiu thue")
    truoc = _tinh_luong(client, emp)
    oid = _don("DH-LUONG-05", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    sau = _tinh_luong(client, emp)
    assert sau["thu_nhap_chiu_thue"] == truoc["thu_nhap_chiu_thue"] + 5_000_000


def test_hoa_hong_la_COT_khong_phai_dong_khoan_nguon_auto(client):
    """Từ 07/09/2026 không còn dòng khoản nguồn `auto`: mọi khoản trên dòng là `employee`/`line`,
    hoa hồng nằm ở cột `hoa_hong`. (Giao diện không còn phải khoá ô của dòng `auto`.)"""
    emp, uid = _sales("nguon auto")
    oid = _don("DH-LUONG-06", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    line = _tinh_luong(client, emp)
    assert line["hoa_hong"] == 5_000_000
    assert all(c["source"] in ("employee", "line") for c in line["components"])


def test_KHONG_CHO_GO_TAY_hoa_hong(client):
    """⭐ Không có ô nào sửa được số hoa hồng (chốt 24/08/2026: *"kệ nó ăn theo đơn hàng cho chắc"*).
    `PUT /lines/{id}` không nhận `hoa_hong`; sửa một ô khác cũng KHÔNG làm hoa hồng đổi hay mất.
    """
    emp, uid = _sales("go tay")
    _hoa_don(_don("DH-LUONG-07", uid, truoc_vat=100_000_000), 108_000_000)
    line = _tinh_luong(client, emp)
    from app.schemas.payroll import LineUpdateIn
    assert "hoa_hong" not in LineUpdateIn.model_fields
    r = client.put(f"/api/luong/lines/{line['id']}", json={"vi_pham": 0}, headers=_h(client))
    assert r.status_code == 200, r.text
    assert r.json()["hoa_hong"] == 5_000_000
    assert r.json()["gross"] == line["gross"]


def test_HOA_HONG_khong_hien_trong_DANH_MUC_khoan_thu_nhap(client):
    """⭐ Chốt của chủ 24/08/2026 (và nhắc lại 07/09): danh mục là chỗ khai phụ cấp/thưởng thêm/xoá
    tự do; hoa hồng "là một dạng lương" — không nằm đó. Từ 07/09 không còn dòng `hoa_hong_kd` nào
    trong DB để mà giấu, và engine không cần dòng đó để chạy."""
    r = client.get("/api/luong/components", headers=_h(client))
    assert r.status_code == 200, r.text
    ma = [c["code"] for c in r.json()["items"]]
    assert "hoa_hong_kd" not in ma
    assert "khoan_km_gh" not in ma, "tàn dư bản nháp mg 0231 — khoán km là cột riêng"
    assert ma, "các khoản phụ cấp/thưởng phải còn nguyên"
    # Vế hai: không có dòng danh mục mà engine VẪN ra tiền.
    emp, uid = _sales("van chay")
    _hoa_don(_don("DH-LUONG-07E", uid, truoc_vat=100_000_000), 108_000_000)
    assert _tinh_luong(client, emp)["hoa_hong"] == 5_000_000


def test_FILE_XUAT_hoa_hong_co_COT_RIENG_khong_lan_vao_Thuong(client):
    """⭐ Hoa hồng nằm trong cột "Tổng" thì phải có MỘT cột mang ĐÚNG TÊN nó giải thích.

    Hai đời luật, cả hai đều từ một lỗi thật:

    1. 21/08/2026 — cột "Thưởng" chỉ cộng `source='line'`, khoản `auto` lọt vào Tổng mà không cột
       nào nói ra. Kế toán dò lệch mãi không ra (y như "Cơm ca"/"Phụ cấp ca" ngày 03/08).
       ⇒ vá bằng cách cộng `auto` vào "Thưởng".
    2. 24/08/2026 — vá kiểu đó thành GIẤU CHỖ KHÁC: chủ mở bảng lương tìm cột hoa hồng không thấy,
       vì nó nằm lẫn với thưởng nóng trong một con số chung, chi tiết chỉ hiện khi rê chuột.
       ⇒ tách hẳn thành cột "Hoa hồng".

    Nên test này khoá HAI chiều: cột Hoa hồng phải nhận đủ tiền, và cột Thưởng phải KHÔNG đổi —
    thiếu vế sau thì cộng hai cột lại vượt "Tổng" mà vẫn xanh.

    Từ 17/09/2026 file theo khuôn `BL CT` của công ty: hoa hồng nằm ở cột "Lương kinh doanh/Sản lượng" (cột Q
    của họ — người kinh doanh không có sản lượng, thợ không có hoa hồng), vẫn là cột mang tên nó, không lẫn thưởng.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    def _xuat(ten_nv: str) -> tuple[list, tuple]:
        r = client.get("/api/luong/export.xlsx?year=2026&month=8", headers=_h(client))
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        # Khuôn `BL CT` (17/09/2026): tiêu đề DÒNG 4, dữ liệu từ dòng 5, cột 3 là Họ và tên.
        head = [c.value for c in ws[4]]
        return head, next(x for x in ws.iter_rows(min_row=5, values_only=True) if x[2] == ten_nv)

    emp, uid = _sales("file xuat")
    _tinh_luong(client, emp)
    head, truoc = _xuat("file xuat")

    oid = _don("DH-LUONG-08", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)
    _tinh_luong(client, emp)
    _, sau = _xuat("file xuat")

    i_thuong, i_hh, i_tong = (head.index("Thưởng/khoản phát sinh"),
                              head.index("Lương kinh doanh/Sản lượng"), head.index("Tổng lương"))
    assert sau[i_tong] - truoc[i_tong] == 5_000_000, "Tổng không nhận hoa hồng"
    assert sau[i_hh] - truoc[i_hh] == 5_000_000, (
        "Tổng có thêm 5tr mà cột Lương kinh doanh không tăng — file xuất cộng lại không khớp")
    assert sau[i_thuong] == truoc[i_thuong], (
        "hoa hồng vẫn còn lẫn trong cột Thưởng ⇒ cộng hai cột lại là đếm đôi")


def test_NV_DA_NGHI_VIEC_van_duoc_tra_hoa_hong_don_cu(client):
    """⭐ Nghỉ việc rồi hoá đơn mới ra ⇒ vẫn còn tiền phải trả, KHÔNG được quỵt.

    Vòng `generate` bỏ qua NV nghỉ việc "không có công/khoán/dòng lương trong kỳ". NV kinh doanh
    nghỉ tháng trước, tháng này kế toán mới xuất hoá đơn của đơn họ đã chốt — không có công, không
    có khoán, chưa có dòng lương nào ⇒ trước bản vá là rơi khỏi bảng lương, tiền bốc hơi không một
    dòng cảnh báo.
    """
    from app.models.employee import STATUS_RESIGNED

    emp, uid = _sales("da nghi viec")
    oid = _don("DH-LUONG-09", uid, truoc_vat=100_000_000)
    _hoa_don(oid, 108_000_000)

    db = SessionLocal()
    try:
        e = db.get(Employee, emp)
        e.status = STATUS_RESIGNED
        e.resign_date = date(2026, 7, 31)
        db.commit()
    finally:
        db.close()

    r = client.post("/api/luong/generate", json={"year": 2026, "month": 8}, headers=_h(client))
    assert r.status_code in (200, 201), r.text
    dong = next((l for l in r.json()["lines"] if l["employee_id"] == emp), None)
    assert dong is not None, "NV nghỉ việc còn hoa hồng mà không có dòng lương nào"
    assert dong["hoa_hong"] == 5_000_000


def test_KHONG_CO_DUONG_NAO_de_sale_tu_dat_pct_cho_don(client):
    """⭐ Ai đặt % hoa hồng — chốt bằng SỐ, không bằng lời (chủ 21/08/2026).

    "% của người kinh doanh nó bán được đơn rồi thì nó theo chỗ mình setup trước đó chứ, làm sao
    nó có quyền chỉnh sửa % hoa hồng được." Cho sale gõ % trên chính đơn mình bán là để người ta
    TỰ VIẾT PHIẾU LƯƠNG CỦA MÌNH.

    Test này ĐỎ ngay khi ai đó phơi `commission_pct` ra API đơn hàng.
    """
    from app.schemas.order import OrderDetailOut, OrderUpdate

    assert "commission_pct" not in OrderUpdate.model_fields, (
        "đã mở đường ghi % hoa hồng qua API sửa đơn — sale tự đặt được hoa hồng của mình")
    assert "commission_pct" not in OrderDetailOut.model_fields, (
        "% hoa hồng của người khác lộ ra trên đơn — ai xem được đơn là biết mức hoa hồng")
    for f in ("bac_hoa_hong", "hh_rieng_pct", "hh_rieng_ly_do"):
        assert f not in OrderUpdate.model_fields and f not in OrderDetailOut.model_fields, f




# =============================================================================================
# Ngăn soi hoa hồng + % riêng của đơn (23/10/2026)
# =============================================================================================
def test_SOI_hoa_hong_tong_khop_cot_dong_luong(client):
    emp, uid = _sales("soi")
    oid = _don("DH-SOI", uid, truoc_vat=100_000_000, bac=CTV)
    iid = _hoa_don_chi(oid, 108_000_000, ngay=date(2026, 7, 5))
    _thu([(oid, iid, 54_000_000)], ngay=date(2026, 8, 10))     # 36 ngày ⇒ 3%
    _thu([(oid, iid, 54_000_000)], ngay=date(2026, 8, 25))     # 51 ngày ⇒ 2%
    dong = _tinh_luong(client, emp)
    assert dong["hoa_hong"] == 1_500_000 + 1_000_000
    r = client.get(f"/api/luong/lines/{dong['id']}/hoa-hong", headers=_h(client))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["tong"] == dong["hoa_hong"]
    assert [(k["so_ngay"], k["pct"], round(k["tien"])) for k in d["items"]] == [
        (36, 0.03, 1_500_000), (51, 0.02, 1_000_000)]
    assert d["items"][0]["order_no"] == "DH-SOI" and d["items"][0]["pct_rieng"] is False


def test_dat_hh_rieng_doi_tien_va_ghi_nhat_ky(client):
    from app.models.audit import AuditLog

    emp, uid = _sales("rieng api")
    oid = _don("DH-RIENG", uid, truoc_vat=100_000_000, bac=CTV)
    _hoa_don(oid, 108_000_000)
    r = client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0.01, "ly_do": "Đơn Wipro"},
                   headers=_h(client))
    assert r.status_code == 200, r.text
    assert r.json() == {"order_id": oid, "pct": 0.01, "ly_do": "Đơn Wipro"}
    assert _tinh(emp) == 1_000_000
    db = SessionLocal()
    try:
        assert db.query(AuditLog).filter(AuditLog.action == "payroll_set_hh_rieng",
                                         AuditLog.target == f"order:{oid}").count() == 1
    finally:
        db.close()


def test_hh_rieng_thieu_ly_do_400(client):
    _emp, uid = _sales("rieng thieu")
    oid = _don("DH-RIENG-LD", uid, truoc_vat=100_000_000)
    r = client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0.01, "ly_do": "  "},
                   headers=_h(client))
    assert r.status_code == 400, r.text


def test_bo_hh_rieng_ve_None_thi_quay_lai_bieu(client):
    emp, uid = _sales("rieng bo")
    oid = _don("DH-RIENG-BO", uid, truoc_vat=100_000_000, bac=CTV)
    _hoa_don(oid, 108_000_000)
    client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0, "ly_do": "Loại tay"},
               headers=_h(client))
    assert _tinh(emp) == 0
    r = client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": None, "ly_do": None},
                   headers=_h(client))
    assert r.status_code == 200, r.text
    assert _tinh(emp) == 3_000_000


def test_SALE_khong_dat_duoc_hh_rieng(client):
    """⭐ Sale không có ô `luong:update` ⇒ 403: không ai tự viết hoa hồng cho đơn mình bán."""
    from tests.test_luong_api import _sales_token

    _emp, uid = _sales("rieng sale")
    oid = _don("DH-RIENG-SALE", uid, truoc_vat=100_000_000)
    r = client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0.5, "ly_do": "tự đặt"},
                   headers={"Authorization": f"Bearer {_sales_token()}"})
    assert r.status_code == 403, r.text


def test_hh_rieng_don_cua_nguoi_ngoai_pham_vi_403(client, monkeypatch):
    """Người có `luong:update` nhưng phạm vi chỉ một phòng không đổi được hoa hồng đơn của người ngoài
    phạm vi — cùng hàng rào `chan_ngoai_pham_vi` với sửa dòng lương."""
    from app.services.payroll_service import PayrollService

    _emp, uid = _sales("rieng ngoai pv")
    oid = _don("DH-RIENG-PV", uid, truoc_vat=100_000_000)
    monkeypatch.setattr(PayrollService, "nv_duoc_xem", lambda self, *, scope, actor: {-1})
    r = client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0.01, "ly_do": "x"},
                   headers=_h(client))
    assert r.status_code == 403, r.text


def test_SOI_tong_kep_0_nhu_cot_bang_luong(client):
    """Người chỉ có một lần thu gốc âm: cột bảng lương kẹp 0 ⇒ tổng ngăn soi cũng 0 (không báo lệch
    oan), từng khoản vẫn hiện số âm để soát."""
    emp, uid = _sales("soi am")
    am = _don("DH-SOI-AM", uid, truoc_vat=86_190_000, vat_pct=8, bac=LANH_DAO)
    iid = _hoa_don_chi(am, 93_085_200, ngay=date(2026, 8, 1), chiet_khau=321_781_872)
    _thu([(am, iid, 93_085_200)], ngay=date(2026, 8, 20))
    dong = _tinh_luong(client, emp)
    assert dong["hoa_hong"] == 0
    d = client.get(f"/api/luong/lines/{dong['id']}/hoa-hong", headers=_h(client)).json()
    assert d["tong"] == 0 and d["items"][0]["tien"] < 0


def test_SOI_tra_ly_do_pct_rieng_de_dien_san(client):
    emp, uid = _sales("soi ly do")
    oid = _don("DH-SOI-LD", uid, truoc_vat=100_000_000, bac=CTV)
    _hoa_don(oid, 108_000_000)
    client.put(f"/api/luong/don/{oid}/hoa-hong-rieng", json={"pct": 0.01, "ly_do": "Đơn Wipro"},
               headers=_h(client))
    dong = _tinh_luong(client, emp)
    k = client.get(f"/api/luong/lines/{dong['id']}/hoa-hong", headers=_h(client)).json()["items"][0]
    assert (k["pct_rieng"], k["pct"], k["ly_do_rieng"]) == (True, 0.01, "Đơn Wipro")


def test_bieu_rong_luu_SQL_NULL_khong_phai_json_null(client):
    """`bac_hoa_hong = None` phải là SQL NULL — JSON 'null' làm `IS NULL` (lọc người không hưởng) trượt."""
    from sqlalchemy import text

    emp, uid = _sales("bieu rong", pct=0, bac=None)
    oid = _don("DH-BIEU-RONG", uid, truoc_vat=1, pct=None)
    db = SessionLocal()
    try:
        assert db.execute(text("SELECT bac_hoa_hong IS NULL FROM employee_salaries WHERE employee_id = :e"),
                          {"e": emp}).scalar() in (1, True)
        assert db.execute(text("SELECT bac_hoa_hong IS NULL FROM orders WHERE id = :o"),
                          {"o": oid}).scalar() in (1, True)
    finally:
        db.close()


def test_CHOT_DON_QUA_API_chup_bieu_cua_nguoi_phu_trach(client):
    """Đường chốt thật (`POST /orders/{id}/confirm`), không gọi thẳng helper."""
    from tests.test_cham_do_don_hang import _bao_gia_da_chot

    bieu = [{"den_ngay": 45, "pct": 0.03}, {"den_ngay": None, "pct": 0}]
    _emp, uid = _sales("chot api", bac=bieu)
    h = _h(client)
    d = client.post("/api/orders", json={"source_type": "bao_gia", "quotation_id": _bao_gia_da_chot(),
                                         "deposit_pct": 0}, headers=h).json()
    client.put(f"/api/orders/{d['id']}", json={"customer_po_no": "PO-HH",
                                               "delivery_committed_date": date.today().isoformat()},
               headers=h)
    db = SessionLocal()
    try:
        db.get(Order, d["id"]).sale_user_id = uid
        db.commit()
    finally:
        db.close()
    r = client.post(f"/api/orders/{d['id']}/confirm", headers=h)
    assert r.status_code == 200, r.text
    db = SessionLocal()
    try:
        assert db.get(Order, d["id"]).bac_hoa_hong == bieu
    finally:
        db.close()


def test_TONG_AM_van_danh_dau_co_lan_thu_de_mo_soi(client):
    """Cột kẹp 0 nhưng người này CÓ lần thu (tổng âm) ⇒ `co_hoa_hong` bật để màn cho bấm soi. Người
    không có lần thu nào thì tắt."""
    emp, uid = _sales("co hh am")
    am = _don("DH-CO-AM", uid, truoc_vat=86_190_000, vat_pct=8, bac=LANH_DAO)
    iid = _hoa_don_chi(am, 93_085_200, ngay=date(2026, 8, 1), chiet_khau=321_781_872)
    _thu([(am, iid, 93_085_200)], ngay=date(2026, 8, 20))
    khong, _u = _sales("co hh khong")
    r = client.post("/api/luong/generate", json={"year": 2026, "month": 8}, headers=_h(client))
    dong = {l["employee_id"]: l for l in r.json()["lines"]}
    assert (dong[emp]["hoa_hong"], dong[emp]["co_hoa_hong"]) == (0, True)
    assert dong[khong]["co_hoa_hong"] is False
