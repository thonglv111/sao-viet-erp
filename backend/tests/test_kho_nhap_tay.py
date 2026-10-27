"""Thêm / cập nhật / xoá / import / export hàng ngay trên màn Tồn kho (spec 2026-10-28).

Đi bằng HTTP thật như màn kho dùng: `/api/kho/ton/{kho_id}/hang/*`.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models.kho_khoa_so import KhoKhoaSo
from app.models.role import SCOPE_ALL
from app.models.stock_lot import StockLot
from app.models.stock_request import StockRequest
from app.models.stock_voucher import StockVoucher, StockVoucherLine
from app.models.vat_lieu_kho import GiayNguyen, VatTuInAn
from app.repositories.rbac_repo import DepartmentRepository, RoleRepository
from app.repositories.user_repo import UserRepository
from app.security import hash_password
from app.services.quyen_kho import khoa_kho
from tests.test_kho_de_nghi import _admin, _login, _mk_kho
from tests.test_kho_lo_giay import _don_vi

PW = "pw123456"


def _nguoi(username: str, kho_id: int, *, adjust: bool = True) -> dict:
    db = SessionLocal()
    try:
        dept = DepartmentRepository(db).get_by_name("Kho") or DepartmentRepository(db).create(name="Kho")
        roles = RoleRepository(db)
        role = roles.create(name=f"Vai {username}", department_id=dept.id)
        roles.set_permission(role_id=role.id, module_key=khoa_kho(kho_id), can_read=True,
                             can_adjust=adjust, can_view_cost=True, scope=SCOPE_ALL)
        u = UserRepository(db).create(username=username, name=username, password_hash=hash_password(PW))
        UserRepository(db).set_assignment(u, department_id=dept.id, role_id=role.id, is_active=True)
        db.commit()
    finally:
        db.close()


def _giay(ma="COUCHE150", gsm=150, them=("to",), don_vi_gia="kg") -> int:
    db = SessionLocal()
    try:
        g = GiayNguyen(ma=ma, ten=f"Giấy {ma}", gsm=gsm, don_vi_gia=don_vi_gia, don_vi_them=list(them))
        db.add(g)
        db.commit()
        return g.id
    finally:
        db.close()


def _vat_tu(ma="KEO", dv="kg") -> int:
    db = SessionLocal()
    try:
        v = VatTuInAn(ma=ma, ten=f"Vật tư {ma}", don_vi_gia=dv)
        db.add(v)
        db.commit()
        return v.id
    finally:
        db.close()


def _moc(client):
    admin = _admin(client)
    kho_id = _mk_kho(client, admin)
    _don_vi()
    _nguoi("thukho", kho_id)
    return kho_id, _login(client, "thukho")


def _them(client, kho_id, h, **d):
    return client.post(f"/api/kho/ton/{kho_id}/hang/them", headers=h, json=d)


def _lo(kho_id):
    db = SessionLocal()
    try:
        return db.query(StockLot).filter(StockLot.kho_id == kho_id).order_by(StockLot.id).all()
    finally:
        db.close()


def _ton(kho_id) -> float:
    return sum(float(x.sl_con_lai) for x in _lo(kho_id))


# ---------- Thêm / cập nhật ----------


def test_them_moi_tao_lo_va_phieu_nhap_tay(client):
    kho_id, h = _moc(client)
    _giay()
    r = _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=1000, don_gia=3150,
              kho_rong=1090, kho_dai=790, ncc="Tân Mai", ngay_nhap="2026-09-30")
    assert r.status_code == 200, r.text
    kq = r.json()
    assert kq["da_ghi"] and kq["tao_moi"] == 1, kq
    (lo,) = _lo(kho_id)
    assert (lo.kho_rong, lo.kho_dai, lo.don_vi_goc) == (790, 1090, "to")
    assert float(lo.sl_con_lai) == 1000 and int(lo.don_gia_nhap) == 3150
    assert lo.ncc == "Tân Mai" and lo.ngay_nhap == date(2026, 9, 30)
    db = SessionLocal()
    try:
        v = db.get(StockVoucher, lo.voucher_id)
        assert v.trang_thai == "posted" and v.ngay == date(2026, 9, 30)
        assert db.get(StockRequest, v.request_id).nhap_tay is True
    finally:
        db.close()


def test_xem_truoc_khong_ghi(client):
    kho_id, h = _moc(client)
    _giay()
    r = _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=10, don_gia=1, kho_rong=790,
              kho_dai=1090, xem_truoc=True)
    assert r.json()["hop_le"] and not r.json()["da_ghi"]
    assert _lo(kho_id) == []


def test_trung_la_cap_nhat_tang_giam_bang(client):
    kho_id, h = _moc(client)
    _giay()
    base = dict(ma="COUCHE150", dvt="tờ", kho_rong=790, kho_dai=1090)
    assert _them(client, kho_id, h, so_luong=500, don_gia=3000, **base).json()["tao_moi"] == 1
    # Tăng: phần tăng là lô mới theo giá mới.
    kq = _them(client, kho_id, h, so_luong=800, don_gia=3300, **base).json()
    assert kq["cap_nhat"] == 1 and kq["thay_doi"][0]["cu"].startswith("500")
    lots = _lo(kho_id)
    assert [float(x.sl_con_lai) for x in lots] == [500, 300]
    assert [int(x.don_gia_nhap) for x in lots] == [3000, 3300]
    # Giảm: trừ lô nhập trước, giá lô không đổi.
    _them(client, kho_id, h, so_luong=600, **base)
    lots = _lo(kho_id)
    assert [float(x.sl_con_lai) for x in lots] == [300, 300]
    assert [int(x.don_gia_nhap) for x in lots] == [3000, 3300]
    # Bằng: không đổi.
    kq = _them(client, kho_id, h, so_luong=600, **base).json()
    assert kq["khong_doi"] == 1 and kq["cap_nhat"] == 0
    # Khác khổ là dòng khác.
    kq = _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", kho_rong=650, kho_dai=860,
               so_luong=10, don_gia=2000).json()
    assert kq["tao_moi"] == 1


def test_tang_so_ngay_cu_hien_trong_xem_truoc(client):
    # File xuất ra mang Ngày nhập của lô gần nhất; tăng số mà giữ ngày đó thì phần tăng lùi về ngày cũ —
    # xem trước phải nói ra, không để ngầm.
    kho_id, h = _moc(client)
    _giay()
    base = dict(ma="COUCHE150", dvt="tờ", kho_rong=790, kho_dai=1090)
    _them(client, kho_id, h, so_luong=500, don_gia=3000, ngay_nhap="2026-09-01", **base)
    kq = _them(client, kho_id, h, so_luong=800, ngay_nhap="2026-09-01", xem_truoc=True, **base).json()
    assert {"cot": "Ngày nhập phần tăng", "moi": "01/09/2026"}.items() <= next(
        t for t in kq["thay_doi"] if t["cot"] == "Ngày nhập phần tăng").items()
    # Ngày trống = hôm nay ⇒ không thêm dòng.
    kq = _them(client, kho_id, h, so_luong=800, xem_truoc=True, **base).json()
    assert all(t["cot"] != "Ngày nhập phần tăng" for t in kq["thay_doi"])


def test_nhan_nhap_tay_o_yeu_cau_phieu_va_bao_cao(client):
    # Spec §7: cờ nhap_tay để danh sách yêu cầu, phiếu và báo cáo kho gọi đúng tên nguồn.
    kho_id, h = _moc(client)
    _giay()
    base = dict(ma="COUCHE150", dvt="tờ", kho_rong=790, kho_dai=1090)
    _them(client, kho_id, h, so_luong=500, don_gia=3000, **base)
    _them(client, kho_id, h, so_luong=200, **base)
    admin = _admin(client)
    ycs = client.get("/api/kho/de-nghi", params={"kho_id": kho_id}, headers=admin).json()["items"]
    assert ycs and all(r["nhap_tay"] for r in ycs)
    ps = client.get("/api/kho/phieu", params={"kho_id": kho_id}, headers=admin).json()["items"]
    assert ps and all(p["nhap_tay"] for p in ps)
    dong = client.get("/api/kho/bao-cao/dong", params={"kho_id": kho_id}, headers=admin).json()["items"]
    assert {(d["ly_do"], d["ly_do_ten"]) for d in dong} == {
        ("nhap_tay", "Nhập tay ở Tồn kho"), ("xuat_tay", "Giảm tay ở Tồn kho")}


def test_ram_quy_ve_to_va_trung_dong(client):
    kho_id, h = _moc(client)
    _giay()
    kq = _them(client, kho_id, h, ma="COUCHE150", dvt="ram", so_luong=2, don_gia=1_500_000,
               kho_rong=790, kho_dai=1090).json()
    assert kq["da_ghi"], kq
    (lo,) = _lo(kho_id)
    assert float(lo.sl_con_lai) == 1000 and int(lo.don_gia_nhap) == 3000
    # 1000 tờ cùng khổ = cùng dòng ⇒ không đổi
    kq = _them(client, kho_id, h, ma="couche150", dvt="Tờ", so_luong=1000, kho_rong=1090,
               kho_dai=790).json()
    assert kq["khong_doi"] == 1


def test_kg_kho_hai_canh_quy_ra_to(client):
    kho_id, h = _moc(client)
    _giay(them=("to", "kg"))
    kq = _them(client, kho_id, h, ma="COUCHE150", dvt="kg", so_luong=129.15, don_gia=21000,
               kho_rong=790, kho_dai=1090).json()
    assert kq["da_ghi"], kq
    (lo,) = _lo(kho_id)
    assert lo.don_vi_goc == "to" and float(lo.sl_con_lai) == 999   # 129,15 ÷ 0,129165 = 999,88 ⇒ làm tròn xuống


def test_loi_dvt(client):
    kho_id, h = _moc(client)
    _giay()
    _giay(ma="NOGSM", gsm=0, them=("to", "kg"))
    base = dict(ma="COUCHE150", so_luong=10, don_gia=1)

    def loi(**d):
        kq = _them(client, kho_id, h, **{**base, **d}).json()
        assert not kq["hop_le"], kq
        return " | ".join(x["ly_do"] for x in kq["loi"])

    assert "Chưa ghi ĐVT" in loi(dvt="", kho_rong=790, kho_dai=1090)
    assert "ream" in loi(dvt="ream", kho_rong=790, kho_dai=1090)
    assert "không đếm theo tờ" in loi(dvt="tờ", kho_rong=650)
    assert "định lượng" in loi(ma="NOGSM", dvt="kg", kho_rong=790, kho_dai=1090)


def test_loi_ma_va_gia(client):
    kho_id, h = _moc(client)
    _giay(ma="TRUNG")
    _vat_tu(ma="TRUNG")
    kq = _them(client, kho_id, h, ma="TRUNG", dvt="kg", so_luong=1, don_gia=1).json()
    assert "cả danh mục giấy lẫn danh mục vật tư" in kq["loi"][0]["ly_do"]
    kq = _them(client, kho_id, h, ma="KHONGCO", dvt="kg", so_luong=1, don_gia=1).json()
    assert "Không có mã" in kq["loi"][0]["ly_do"]
    _vat_tu(ma="KEO")
    kq = _them(client, kho_id, h, ma="KEO", dvt="kg", so_luong=1).json()
    assert any("phải ghi đơn giá" in x["ly_do"] for x in kq["loi"]), kq
    kq = _them(client, kho_id, h, ma="KEO", dvt="kg", so_luong=-1, don_gia=1).json()
    assert any("không được âm" in x["ly_do"] for x in kq["loi"]), kq


def test_ky_khoa_chan_ngay_nhap(client):
    kho_id, h = _moc(client)
    _vat_tu(ma="KEO")
    db = SessionLocal()
    try:
        db.add(KhoKhoaSo(kho_id=None, tu_ngay=date(2026, 8, 1), den_ngay=date(2026, 8, 31), hanh_dong="khoa"))
        db.commit()
    finally:
        db.close()
    kq = _them(client, kho_id, h, ma="KEO", dvt="kg", so_luong=1, don_gia=1, ngay_nhap="15/08/2026").json()
    assert any("khoá sổ" in x["ly_do"] for x in kq["loi"]), kq


def test_vat_tu_co_kho_tuy_chon(client):
    kho_id, h = _moc(client)
    _vat_tu(ma="BOPP", dv="m")
    assert _them(client, kho_id, h, ma="BOPP", dvt="m", so_luong=100, don_gia=500, kho_rong=320).json()["da_ghi"]
    assert _them(client, kho_id, h, ma="BOPP", dvt="m", so_luong=50, don_gia=520, kho_rong=350).json()["tao_moi"] == 1
    assert {(x.kho_rong, x.kho_dai) for x in _lo(kho_id)} == {(320, 0), (350, 0)}


# ---------- quyền ----------


def test_khong_co_quyen_403(client):
    admin = _admin(client)
    kho_id = _mk_kho(client, admin)
    _nguoi("xemthoi", kho_id, adjust=False)
    h = _login(client, "xemthoi")
    assert _them(client, kho_id, h, ma="X").status_code == 403
    assert client.get(f"/api/kho/ton/{kho_id}/hang/xuat-excel", headers=h).status_code == 403
    assert client.post(f"/api/kho/ton/{kho_id}/hang/xoa", headers=h, json={"dong": [
        {"hang_loai": "giay", "hang_id": 1}]}).status_code == 403


def test_ma_tran_bat_adjust_keo_xem_gia(client):
    admin = _admin(client)
    kho_id = _mk_kho(client, admin)
    db = SessionLocal()
    try:
        dept = DepartmentRepository(db).get_by_name("Kho") or DepartmentRepository(db).create(name="Kho")
        role = RoleRepository(db).create(name="Vai thử ma trận", department_id=dept.id)
        db.commit()
        rid = role.id
    finally:
        db.close()
    r = client.put(f"/api/roles/{rid}/permissions", headers=admin, json={"permissions": [
        {"module_key": khoa_kho(kho_id), "can_read": True, "can_adjust": True, "scope": "all"}]})
    assert r.status_code in (200, 204), r.text
    db = SessionLocal()
    try:
        p = RoleRepository(db).get_permission(rid, khoa_kho(kho_id))
        assert p.can_adjust and p.can_view_cost
    finally:
        db.close()


# ---------- xoá ----------


def test_xoa_xoa_han_lo_va_phieu_nhap_tay(client):
    kho_id, h = _moc(client)
    gid = _giay()
    _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=500, don_gia=3000, kho_rong=790, kho_dai=1090)
    _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=300, kho_rong=790, kho_dai=1090)
    assert len(_lo(kho_id)) == 1
    r = client.post(f"/api/kho/ton/{kho_id}/hang/xoa", headers=h, json={"dong": [
        {"hang_loai": "giay", "hang_id": gid, "kho_rong": 1090, "kho_dai": 790, "don_vi_goc": "to"}]})
    assert r.status_code == 200, r.text
    assert _lo(kho_id) == []
    db = SessionLocal()
    try:
        assert db.query(StockVoucherLine).count() == 0
        assert db.query(StockVoucher).count() == 0
        assert db.query(StockRequest).filter(StockRequest.nhap_tay.is_(True)).count() == 0
    finally:
        db.close()


# ---------- Excel ----------


def _file(client, kho_id, h, url="mau-excel"):
    r = client.get(f"/api/kho/ton/{kho_id}/hang/{url}", headers=h)
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))


def _gui(client, kho_id, h, wb, mode="preview"):
    buf = BytesIO()
    wb.save(buf)
    return client.post(f"/api/kho/ton/{kho_id}/hang/import-excel?mode={mode}", headers=h,
                       files={"file": ("ton.xlsx", buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _ghi_dong(ws, r, **o):
    nhan = [c.value for c in ws[1]]
    for k, v in o.items():
        col = next(i for i, n in enumerate(nhan, start=1) if n.replace(" (*)", "") == k)
        ws.cell(row=r, column=col, value=v)


def test_import_xem_truoc_roi_ghi(client):
    kho_id, h = _moc(client)
    _giay()
    _vat_tu(ma="KEO")
    wb = _file(client, kho_id, h)
    ws = wb["Tồn kho"]
    _ghi_dong(ws, 2, **{"Mã hàng": "COUCHE150", "Khổ rộng": 790, "Khổ dài": 1090, "ĐVT": "ram",
                        "Số lượng": 2, "Đơn giá": 1_500_000})
    _ghi_dong(ws, 3, **{"Mã hàng": "KEO", "ĐVT": "kg", "Số lượng": 18, "Đơn giá": 185000,
                        "Ngày nhập": "01/09/2026"})
    kq = _gui(client, kho_id, h, wb).json()
    assert kq["hop_le"] and kq["tao_moi"] == 2 and not kq["da_ghi"], kq
    assert any("= 1.000 tờ" in t["moi"] for t in kq["thay_doi"]), kq["thay_doi"]
    assert _lo(kho_id) == []
    kq = _gui(client, kho_id, h, wb, "commit").json()
    assert kq["da_ghi"] and kq["tao_moi"] == 2
    assert _ton(kho_id) == 1018


def test_import_trung_ram_va_to(client):
    kho_id, h = _moc(client)
    _giay()
    wb = _file(client, kho_id, h)
    ws = wb["Tồn kho"]
    for r, (dvt, sl) in enumerate((("ram", 2), ("tờ", 1000)), start=2):
        _ghi_dong(ws, r, **{"Mã hàng": "COUCHE150", "Khổ rộng": 790, "Khổ dài": 1090, "ĐVT": dvt,
                            "Số lượng": sl, "Đơn giá": 1})
    kq = _gui(client, kho_id, h, wb).json()
    assert not kq["hop_le"] and {x["dong"] for x in kq["loi"]} == {2, 3}, kq


def test_import_loi_mot_dong_khong_ghi_gi(client):
    kho_id, h = _moc(client)
    _vat_tu(ma="KEO")
    wb = _file(client, kho_id, h)
    ws = wb["Tồn kho"]
    _ghi_dong(ws, 2, **{"Mã hàng": "KEO", "ĐVT": "kg", "Số lượng": 5, "Đơn giá": 1})
    _ghi_dong(ws, 3, **{"Mã hàng": "SAI", "ĐVT": "kg", "Số lượng": 5, "Đơn giá": 1})
    kq = _gui(client, kho_id, h, wb, "commit").json()
    assert not kq["da_ghi"] and kq["loi"][0]["dong"] == 3
    assert _lo(kho_id) == []


def test_import_file_kho_khac_va_thieu_cot(client):
    kho_id, h = _moc(client)
    kho2 = _mk_kho(client, _admin(client))
    _nguoi("thukho2", kho2)
    wb = _file(client, kho_id, h)
    r = _gui(client, kho2, _login(client, "thukho2"), wb)
    assert r.status_code == 422 and "đúng kho" in r.json()["detail"]
    ws = wb["Tồn kho"]
    ws.delete_cols(6)   # bỏ cột ĐVT
    r = _gui(client, kho_id, h, wb)
    assert r.status_code == 422 and "ĐVT" in r.json()["detail"]


def test_export_roi_import_lai_khong_doi(client):
    kho_id, h = _moc(client)
    _giay()
    _vat_tu(ma="BOPP", dv="m")
    _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=493, don_gia=3152, kho_rong=790, kho_dai=1090)
    _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=725, don_gia=3155, kho_rong=790, kho_dai=1090)
    _them(client, kho_id, h, ma="BOPP", dvt="m", so_luong=4000, don_gia=520, kho_rong=320)
    wb = _file(client, kho_id, h, "xuat-excel")
    ws = wb["Tồn kho"]
    assert ws.max_row >= 3
    kq = _gui(client, kho_id, h, wb).json()
    assert kq["hop_le"] and kq["khong_doi"] == 2 and kq["tao_moi"] == kq["cap_nhat"] == 0, kq


# ---------- xem theo ĐVT ----------


def test_ton_nhom_tra_don_vi_xem(client):
    kho_id, h = _moc(client)
    _giay(them=("to", "kg"))
    _them(client, kho_id, h, ma="COUCHE150", dvt="tờ", so_luong=1000, don_gia=3000, kho_rong=790, kho_dai=1090)
    r = client.get(f"/api/kho/phieu/lo/ton-nhom?kho_id={kho_id}", headers=h)
    assert r.status_code == 200, r.text
    (row,) = r.json()["items"]
    dv = {d["ma"]: d["he_so"] for d in row["don_vi_xem"]}
    assert dv["to"] == 1 and dv["ram"] == 500
    assert abs(dv["kg"] - 1 / (0.79 * 1.09 * 0.15)) < 1e-6


def test_o_so_nhan_cham_nghin_kieu_viet():
    from app.services.kho_nhap_tay import _so

    assert _so("1.500.000") == (1_500_000, True)
    assert _so("1.500") == (1500, True)
    assert _so("1.200,5") == (1200.5, True)
    assert _so("2.5") == (2.5, True)
    assert _so("12.34.5") == (None, False)
