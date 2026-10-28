"""Thêm / cập nhật / xoá hàng NGAY TRÊN MÀN TỒN KHO của một kho + nhập / xuất Excel (28/10/2026).

Spec `docs/superpowers/specs/2026-10-28-kho-them-xoa-import-export.md`.

Một DÒNG = một mặt hàng trong kho, khoá bằng **mã + khổ + đơn vị gốc** (đúng dòng của màn Tồn kho).
Form Thêm và file Import đi CÙNG một đường (`doc_dong` → `ap`): trùng dòng đang có ⇒ CẬP NHẬT về đúng
số người ghi, không đẻ dòng mới.

Tiền + đầu kỳ: màn Tồn kho đọc LÔ, báo cáo kho (N-X-T, đầu kỳ, giá BQ) đọc DÒNG PHIẾU. Nên mọi thay
đổi số lượng đều đi qua một cặp yêu cầu + phiếu tự sinh (`stock_requests.nhap_tay`, mg 0451) rồi ghi
sổ bằng chính `StockVoucherService._apply_post` — tạo lô, chặn kỳ đã khoá / đã tính giá y như phiếu
thường. Tăng ⇒ phiếu nhập một lô mới theo đơn giá + ngày nhập người ghi; giảm ⇒ phiếu xuất trừ lô
FEFO rồi FIFO, theo giá của từng lô. Giá lô cũ KHÔNG bao giờ sửa ngược (sửa là đổi giá vốn đã xuất).

XOÁ thì xoá hẳn (chủ chốt 28/10: "xóa thì xóa thôi, đừng chia trường hợp"): mọi lô của dòng trong kho
này, mọi dòng phiếu trỏ tới các lô đó, mức tồn của dòng. Không ghi tăng/giảm. Danh mục không đụng.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field as dc_field
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import delete, select, update
from sqlalchemy.orm import Session

from ..models.document_sequence import (
    SEQ_DOC_TYPE_STOCK_REQUEST_IN,
    SEQ_DOC_TYPE_STOCK_REQUEST_OUT,
    SEQ_DOC_TYPE_STOCK_VOUCHER_IN,
    SEQ_DOC_TYPE_STOCK_VOUCHER_OUT,
)
from ..models.kho_hang import KhoHang, KhoViTri
from ..models.stock_lot import StockLot, StockThreshold
from ..models.stock_request import REQ_APPROVED, REQ_DONE, StockRequest, StockRequestLine
from ..models.stock_voucher import StockVoucher, StockVoucherLine
from ..repositories.audit_repo import AuditLogRepository
from ..repositories.kho_khoa_so_repo import KhoKhoaSoRepository
from ..repositories.kho_ky_ton_repo import KhoKyTonRepository
from .catalog_excel import PHIEN_BAN, SHEET_META, ExcelSaiMan
from .kho_giay import DON_VI_TO_KHO, chuan_kho, kg_tu_to, la_kho_to
from .vat_lieu_kho_service import VatLieuKhoError

LOAI_META = "ton_kho"
SHEET = "Tồn kho"
SHEET_DS = "_ds"
SHEET_HD = "Hướng dẫn"
EPS = 1e-6


class KhoNhapTayError(Exception):
    pass


# ======================= cột file =======================


@dataclass(frozen=True)
class Cot:
    khoa: str
    nhan: str
    rong: int
    bat_buoc: bool
    dinh_dang: str
    can: str
    cach_ghi: str
    tien: bool = False


COT: tuple[Cot, ...] = (
    Cot("ma", "Mã hàng", 13, True, "@", "left", "Mã trong danh mục giấy hoặc danh mục vật tư."),
    Cot("ten", "Tên hàng", 28, False, "@", "left", "Hệ thống lấy theo mã."),
    Cot("dinh_luong", "Định lượng", 10, False, "0", "right", "Chỉ giấy, đơn vị g/m2. Hệ thống lấy theo mã."),
    Cot("kho_rong", "Khổ rộng", 9, False, "0", "right",
        "Đơn vị mm. Hàng không có khổ thì bỏ trống cả khổ rộng và khổ dài."),
    Cot("kho_dai", "Khổ dài", 9, False, "0", "right", "Đơn vị mm. Giấy cuộn, màng chỉ ghi khổ rộng."),
    Cot("dvt", "ĐVT", 8, True, "@", "center",
        "Đơn vị đã khai của mặt hàng, hoặc đơn vị quy đổi được (ví dụ ram, hệ thống quy về tờ)."),
    Cot("so_luong", "Số lượng", 12, True, "General", "right", "Số tồn thực tế theo ĐVT ở cột trước."),
    Cot("don_gia", "Đơn giá", 12, True, "General", "right",
        "Giá theo ĐVT ở cột trước, đơn vị đồng. File xuất ra là giá bình quân của hàng còn tồn.", tien=True),
    Cot("thanh_tien", "Thành tiền", 14, False, "#,##0", "right", "Số lượng nhân đơn giá.", tien=True),
    Cot("vi_tri", "Vị trí", 9, False, "@", "center", "Vị trí đã khai trong kho."),
    Cot("hsd", "Hạn dùng", 11, False, "dd/mm/yyyy", "center", "Dạng dd/mm/yyyy."),
    Cot("ngay_nhap", "Ngày nhập", 11, False, "dd/mm/yyyy", "center",
        "Dạng dd/mm/yyyy. Bỏ trống thì lấy ngày hôm nay."),
    Cot("ncc", "Nhà cung cấp", 20, False, "@", "left", ""),
    Cot("ghi_chu", "Ghi chú", 24, False, "@", "left", ""),
)
NHAN = {c.khoa: c.nhan for c in COT}
LUU_Y = (
    "Mỗi mặt hàng được xác định bằng Mã hàng, Khổ và ĐVT. Mặt hàng đã có trong kho thì dòng trong file sẽ "
    "cập nhật lại số lượng, không tạo dòng mới.",
    "Cùng mã nhưng khác khổ là hai mặt hàng khác nhau.",
    "Mỗi mặt hàng chỉ ghi một dòng. Ghi 2 ram và 1000 tờ cho cùng một khổ là trùng.",
    "Số lượng tăng thì phần tăng thành lô mới theo đơn giá và ngày nhập của dòng. Số lượng giảm thì trừ vào "
    "lô nhập trước.",
    "Ngày nhập trước kỳ báo cáo thì được tính vào tồn đầu kỳ, trong kỳ thì tính vào nhập trong kỳ. Không nhập "
    "được vào kỳ đã khoá sổ hoặc đã tính giá.",
    "File có lỗi thì không dòng nào được ghi. Sửa hết lỗi theo thông báo rồi nhập lại.",
    "File chỉ nhập được vào đúng kho đã tải file về.",
)


def _nhan_goc(nhan: str) -> str:
    """Tiêu đề trong file mang "(*)" ở cột bắt buộc — so khớp thì bỏ dấu đó đi."""
    return nhan.replace("(*)", "").strip().lower()


# ======================= kết quả =======================


@dataclass
class LoiDong:
    dong: int
    cot: str
    ly_do: str


@dataclass
class ThayDoi:
    dong: int
    ma: str
    ten: str
    cot: str
    cu: str
    moi: str


@dataclass
class CanhBao:
    dong: int
    noi_dung: str


@dataclass
class KetQua:
    tong_dong: int = 0
    tao_moi: int = 0
    cap_nhat: int = 0
    khong_doi: int = 0
    da_ghi: bool = False
    loi: list[LoiDong] = dc_field(default_factory=list)
    thay_doi: list[ThayDoi] = dc_field(default_factory=list)
    canh_bao: list[CanhBao] = dc_field(default_factory=list)

    @property
    def hop_le(self) -> bool:
        return not self.loi


@dataclass
class DongVao:
    """Một dòng người ghi — từ form Thêm hoặc từ một dòng file. Ô nào cũng có thể là chữ thô."""

    dong: int = 0
    ma: Any = None
    dvt: Any = None
    so_luong: Any = None
    don_gia: Any = None
    kho_rong: Any = None
    kho_dai: Any = None
    vi_tri: Any = None
    hsd: Any = None
    ngay_nhap: Any = None
    ncc: Any = None
    ghi_chu: Any = None


@dataclass
class KeHoach:
    dong: int
    hang_loai: str
    hang_id: int
    ma: str
    ten: str
    kho_rong: int
    kho_dai: int
    goc: str              # đơn vị gốc (mã) — lô lưu theo đơn vị này
    goc_ten: str
    dvt_ten: str          # đơn vị người ghi (tên)
    dvt: str              # đơn vị người ghi (mã)
    he_so: float          # 1 đơn vị người ghi = he_so đơn vị gốc
    so_luong: float       # theo đơn vị người ghi
    sl_goc: float         # số đích theo đơn vị gốc
    gia_goc: int | None   # đơn giá / đơn vị gốc (None = không ghi)
    don_gia: float | None  # đơn giá theo đơn vị người ghi
    vi_tri: str | None
    hsd: date | None
    ngay: date
    ncc: str | None
    ghi_chu: str | None
    sl_hien: float = 0.0
    co_dong: bool = False
    vi_tri_hien: str | None = None
    lots: list = dc_field(default_factory=list)

    @property
    def khoa(self) -> tuple:
        return (self.hang_loai, self.hang_id, self.kho_rong, self.kho_dai, self.goc)

    @property
    def doi_vi_tri(self) -> bool:
        return self.vi_tri is not None and (self.vi_tri or "") != (self.vi_tri_hien or "")

    @property
    def doi_so(self) -> bool:
        return abs(self.sl_goc - self.sl_hien) > EPS


# ======================= đọc ô =======================


def _chu(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _so(v: Any) -> tuple[float | None, bool]:
    """(số, hợp lệ). Ô trống ⇒ (None, True). Nhận "1.200,5" kiểu Việt lẫn 1200.5."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, bool):
        return None, False
    if isinstance(v, (int, float)):
        return float(v), True
    s = str(v).strip().replace(" ", "")
    # Kiểu Việt: dấu chấm là phân nghìn ("1.500.000", "1.500"), dấu phẩy là thập phân. Không có phẩy
    # mà chấm tách đúng từng nhóm 3 chữ số ⇒ chấm nghìn; còn lại ("2.5") là số thập phân thường.
    if "," in s or re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s), True
    except ValueError:
        return None, False


def _ngay(v: Any) -> tuple[date | None, bool]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, datetime):
        return v.date(), True
    if isinstance(v, date):
        return v, True
    s = str(v).strip()
    for dinh_dang in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, dinh_dang).date(), True
        except ValueError:
            continue
    return None, False


def _hien_so(v: float | None) -> str:
    if v is None:
        return ""
    if abs(v - round(v)) < EPS:
        return f"{int(round(v)):,}".replace(",", ".")
    s = f"{v:,.3f}".rstrip("0").rstrip(".")
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


# ======================= service =======================


class KhoNhapTayService:
    """`vsvc`: `StockVoucherService` — cần `.hang` (tra danh mục + quy đổi), `.sequence`,
    `.lots` và `_apply_post`. Dựng KHÔNG kèm giữ chỗ (xem router): giảm tồn ở đây là đặt lại số
    đếm được, không phải xuất cho lệnh, nên không chặn vì phần đang giữ — giữ chỗ tính lại sau commit."""

    def __init__(self, db: Session, vsvc) -> None:
        self.db = db
        self.v = vsvc
        self.hang = vsvc.hang
        self._dv_ten: dict[str, str] | None = None

    # ---------- tra cứu ----------

    def kho(self, kho_id: int) -> KhoHang:
        k = self.db.get(KhoHang, kho_id)
        if k is None:
            raise KhoNhapTayError("Không tìm thấy kho.")
        return k

    def ten_dv(self, ma: str | None) -> str:
        if self._dv_ten is None:
            self._dv_ten = {d.ma: d.ten for d in self.hang.don_vi.all_rows()}
        return self._dv_ten.get(ma or "", ma or "")

    def vi_tri_kho(self, kho_id: int) -> list[str]:
        return list(self.db.execute(
            select(KhoViTri.ma).where(KhoViTri.kho_id == kho_id, KhoViTri.active.is_(True))
            .order_by(KhoViTri.ma)).scalars())

    def tra_ma(self, ma: str):
        """(hang_loai, bản ghi) theo mã — tra cả Giấy lẫn Vật tư khác (không thành phẩm). Lỗi khi không
        có, hoặc mã có ở CẢ HAI danh mục (mã chỉ duy nhất trong từng danh mục)."""
        tim = []
        g = self.hang.repo.find_by_ma("giay", ma)
        if g is not None:
            tim.append(("giay", g))
        vt = self.hang.repo.find_by_ma("vat_tu", ma)
        if vt is not None and not getattr(vt, "la_thanh_pham", False):
            tim.append(("vat_tu", vt))
        if not tim:
            raise KhoNhapTayError(f"Không có mã “{ma}” trong danh mục giấy hoặc danh mục vật tư.")
        if len(tim) > 1:
            raise KhoNhapTayError(
                f"Mã “{ma}” có ở cả danh mục giấy lẫn danh mục vật tư — đổi mã ở một bên rồi nhập lại.")
        return tim[0]

    def lo_cua_dong(self, kho_id: int, khoa: tuple) -> list[StockLot]:
        """Mọi lô của dòng trong kho (kể cả lô đã hết) — FEFO rồi FIFO, như gợi ý xuất."""
        loai, hid, kr, kd, goc = khoa
        return list(self.db.execute(
            select(StockLot).where(
                StockLot.kho_id == kho_id, StockLot.hang_loai == loai, StockLot.hang_id == hid,
                StockLot.kho_rong == kr, StockLot.kho_dai == kd, StockLot.don_vi_goc == goc,
            ).order_by(StockLot.hsd.is_(None).asc(), StockLot.hsd.asc(),
                       StockLot.ngay_nhap.asc(), StockLot.id.asc())
        ).scalars())

    def _chan_ky(self, kho_id: int, ngay: date) -> str | None:
        if KhoKhoaSoRepository(self.db).is_locked(kho_id, ngay):
            return f"Ngày {ngay:%d/%m/%Y} thuộc kỳ đã khoá sổ."
        ky = KhoKyTonRepository(self.db).ky_da_tinh_tu(kho_id, ngay)
        if ky is not None:
            return (f"Ngày {ngay:%d/%m/%Y} thuộc kỳ đã tính giá "
                    f"({ky[0]:%d/%m/%Y}–{ky[1]:%d/%m/%Y}).")
        return None

    # ---------- đọc một dòng ----------

    def doc_dong(self, kho_id: int, d: DongVao, kq: KetQua, *, vi_tri_hop: set[str] | None,
                 hom_nay: date | None = None) -> KeHoach | None:
        """Kiểm + quy một dòng. Lỗi ghi vào `kq.loi` (dòng, cột, lý do) và trả None."""
        loi_truoc = len(kq.loi)

        def loi(cot: str, ly_do: str) -> None:
            kq.loi.append(LoiDong(d.dong, NHAN[cot], ly_do))

        ma = _chu(d.ma)
        hang = None
        if not ma:
            loi("ma", "Chưa ghi mã hàng.")
        else:
            try:
                hang = self.tra_ma(ma)
            except KhoNhapTayError as e:
                loi("ma", str(e))

        kr_raw, ok_r = _so(d.kho_rong)
        kd_raw, ok_d = _so(d.kho_dai)
        if not ok_r or (kr_raw is not None and kr_raw < 0):
            loi("kho_rong", "Khổ rộng phải là số mm không âm.")
        if not ok_d or (kd_raw is not None and kd_raw < 0):
            loi("kho_dai", "Khổ dài phải là số mm không âm.")
        if kd_raw and not kr_raw:
            loi("kho_rong", "Có khổ dài thì phải ghi cả khổ rộng.")
        kr, kd = chuan_kho(kr_raw or 0, kd_raw or 0)

        sl, ok = _so(d.so_luong)
        if sl is None and ok:
            loi("so_luong", "Chưa ghi số lượng.")
        elif not ok:
            loi("so_luong", "Số lượng phải là số.")
        elif sl < 0:
            loi("so_luong", "Số lượng không được âm.")

        gia, ok = _so(d.don_gia)
        if not ok:
            loi("don_gia", "Đơn giá phải là số.")
        elif gia is not None and gia < 0:
            loi("don_gia", "Đơn giá không được âm.")

        vi_tri = _chu(d.vi_tri)
        if vi_tri and vi_tri_hop is not None:
            khop = {v.lower(): v for v in vi_tri_hop}
            if vi_tri.lower() not in khop:
                loi("vi_tri", f"Kho chưa khai vị trí “{vi_tri}”. Vị trí có: {', '.join(sorted(vi_tri_hop))}.")
            else:
                vi_tri = khop[vi_tri.lower()]

        hsd, ok = _ngay(d.hsd)
        if not ok:
            loi("hsd", "Hạn dùng phải là ngày dạng dd/mm/yyyy.")
        ngay, ok = _ngay(d.ngay_nhap)
        if not ok:
            loi("ngay_nhap", "Ngày nhập phải là ngày dạng dd/mm/yyyy.")
        ngay = ngay or hom_nay or date.today()

        dvt = _chu(d.dvt)
        if not dvt:
            loi("dvt", "Chưa ghi ĐVT.")

        if hang is None or not dvt:
            return None
        loai, obj = hang
        quy = self._quy_dvt(loai, obj, dvt, sl if (sl and sl > 0) else None, kr, kd)
        if isinstance(quy, str):
            loi("dvt", quy)
            return None
        if len(kq.loi) > loi_truoc:
            return None
        goc, goc_ten, he_so, sl_goc, dvt_ma = quy
        if sl == 0:
            sl_goc = 0.0
        gia_goc = None
        if gia is not None:
            gia_goc = round(gia * sl / sl_goc) if sl and sl_goc else round(gia / he_so if he_so else gia)
        return KeHoach(
            dong=d.dong, hang_loai=loai, hang_id=obj.id, ma=obj.ma, ten=obj.ten,
            kho_rong=kr, kho_dai=kd, goc=goc, goc_ten=goc_ten, dvt_ten=self.ten_dv(dvt.lower()) or dvt, dvt=dvt_ma,
            he_so=he_so, so_luong=float(sl), sl_goc=float(sl_goc), gia_goc=gia_goc,
            don_gia=gia, vi_tri=vi_tri, hsd=hsd, ngay=ngay, ncc=_chu(d.ncc), ghi_chu=_chu(d.ghi_chu),
        )

    def _quy_dvt(self, loai: str, obj, dvt: str, sl: float | None, kr: int, kd: int):
        """(gốc, tên gốc, hệ số 1 đvt → gốc, số gốc, mã đvt) — hoặc câu lỗi. Một luật với kho / mua hàng
        (`VatLieuKhoService.quy_ve_goc`), cộng luật riêng của giấy cuộn (khổ chỉ-rộng lưu kg)."""
        thu = sl if sl else 1.0
        try:
            ra = self.hang.quy_ve_goc(loai, obj.id, dvt, thu, kho_rong=kr, kho_dai=kd)
        except VatLieuKhoError as e:
            msg = str(e)
            if "quy ra 0 tờ" in msg and sl:
                return f"{_hien_so(sl)} {dvt} quy ra 0 tờ — số quá nhỏ."
            if "không quy được kg ra tờ" in msg:
                return (f"{obj.ma} chưa có định lượng nên không quy kg ra tờ — ghi theo tờ/ram hoặc khai "
                        f"định lượng ở danh mục.")
            return msg
        goc = (ra.get("don_vi_goc") or "").lower()
        if loai == "giay" and goc == DON_VI_TO_KHO and kr and not la_kho_to(kr, kd):
            return f"Giấy cuộn khổ {kr} không đếm theo tờ — ghi kg hoặc tấn."
        he_so = float(ra["sl_goc"]) / thu if thu else float(ra.get("he_so_ve_goc") or 1)
        return (goc, ra.get("don_vi_goc_ten") or self.ten_dv(goc), he_so, float(ra["sl_goc"]),
                (ra.get("ma_don_vi") or dvt).strip().lower())

    # ---------- so với kho ----------

    def so_voi_kho(self, kho_id: int, k: KeHoach, kq: KetQua) -> None:
        lots = self.lo_cua_dong(kho_id, k.khoa)
        k.lots = lots
        k.co_dong = bool(lots)
        con = [x for x in lots if float(x.sl_con_lai) > EPS]
        k.sl_hien = sum(float(x.sl_con_lai) for x in con)
        moi_nhat = max(lots, key=lambda x: (x.ngay_nhap, x.id)) if lots else None
        k.vi_tri_hien = moi_nhat.vi_tri if moi_nhat is not None else None

    def kiem_sau_so(self, kho_id: int, k: KeHoach, kq: KetQua) -> bool:
        """Luật cần biết dòng đang có: đơn giá bắt buộc với dòng mới, kỳ khoá theo ngày nhập."""
        ok = True
        if not k.co_dong and k.gia_goc is None and k.sl_goc > EPS:
            kq.loi.append(LoiDong(k.dong, NHAN["don_gia"], "Dòng mới phải ghi đơn giá."))
            ok = False
        if k.doi_so:
            chan = self._chan_ky(kho_id, k.ngay)
            if chan:
                kq.loi.append(LoiDong(k.dong, NHAN["ngay_nhap"], chan + " Đổi ngày nhập sang kỳ đang mở."))
                ok = False
        return ok

    def ghi_thay_doi(self, k: KeHoach, kq: KetQua) -> None:
        def them(cot: str, cu: str, moi: str) -> None:
            kq.thay_doi.append(ThayDoi(k.dong, k.ma, k.ten, cot, cu, moi))

        quy = k.dvt_ten.lower() != (k.goc_ten or "").lower()
        if not k.co_dong:
            them("Số lượng", "", f"{_hien_so(k.so_luong)} {k.dvt_ten}"
                 + (f" (= {_hien_so(k.sl_goc)} {k.goc_ten})" if quy else ""))
            if k.don_gia is not None:
                them("Đơn giá", "", f"{_hien_so(k.don_gia)} đ/{k.dvt_ten}"
                     + (f" (= {_hien_so(k.gia_goc)} đ/{k.goc_ten})" if quy else ""))
            return
        if k.doi_so:
            them("Số lượng", f"{_hien_so(k.sl_hien)} {k.goc_ten}",
                 f"{_hien_so(k.so_luong)} {k.dvt_ten}" + (f" (= {_hien_so(k.sl_goc)} {k.goc_ten})" if quy else ""))
            # File xuất mang ngày của lô gần nhất: phần tăng giữ ngày đó là lùi về quá khứ ⇒ nói ra.
            if k.sl_goc > k.sl_hien + EPS and k.ngay != date.today():
                them("Ngày nhập phần tăng", "", k.ngay.strftime("%d/%m/%Y"))
        if k.doi_vi_tri:
            them("Vị trí", k.vi_tri_hien or "", k.vi_tri or "")

    # ---------- ghi ----------

    def _yeu_cau(self, *, user, kho_id: int, loai: str, k: KeHoach, dvt: str, sl: float,
                 gia: int | None) -> StockRequest:
        req = StockRequest(
            ma=self.v.sequence.generate_flat_code(
                SEQ_DOC_TYPE_STOCK_REQUEST_IN if loai == "NHAP" else SEQ_DOC_TYPE_STOCK_REQUEST_OUT),
            loai=loai, nguoi_tao_id=user.id, kho_id=kho_id, trang_thai=REQ_APPROVED, nhap_tay=True,
            nguoi_duyet_id=user.id, duyet_luc=datetime.now(timezone.utc),
            ghi_chu="Nhập tay trên màn Tồn kho" if loai == "NHAP" else "Giảm tồn trên màn Tồn kho",
        )
        req.lines.append(StockRequestLine(
            hang_loai=k.hang_loai, hang_id=k.hang_id, dvt=dvt, sl_de_nghi=sl, sl_duyet=sl,
            sl_da_ung=0, don_gia=gia, don_vi_goc=k.goc, kho_rong=k.kho_rong, kho_dai=k.kho_dai,
            ghi_chu=k.ghi_chu,
        ))
        self.db.add(req)
        self.db.flush()
        return req

    def _phieu(self, *, user, kho_id: int, req: StockRequest, ngay: date, dong: list[dict]) -> StockVoucher:
        v = StockVoucher(
            ma=self.v.sequence.generate_flat_code(
                SEQ_DOC_TYPE_STOCK_VOUCHER_IN if req.loai == "NHAP" else SEQ_DOC_TYPE_STOCK_VOUCHER_OUT),
            loai=req.loai, request_id=req.id, kho_id=kho_id, ngay=ngay, nguoi_lap_id=user.id,
            nguoi_giao_nhan=getattr(user, "name", None), ghi_chu=req.ghi_chu,
        )
        for d in dong:
            v.lines.append(StockVoucherLine(request_line_id=req.lines[0].id, **d))
        self.db.add(v)
        self.db.flush()
        return v

    def ap(self, kho_id: int, k: KeHoach, user) -> None:
        """Đặt dòng về đúng `k.sl_goc` (+ vị trí). KHÔNG commit."""
        if k.sl_goc > k.sl_hien + EPS:
            tang = k.sl_goc - k.sl_hien
            dvt, sl, gia = self._dong_tang(k, tang)
            req = self._yeu_cau(user=user, kho_id=kho_id, loai="NHAP", k=k, dvt=dvt, sl=sl, gia=gia)
            v = self._phieu(user=user, kho_id=kho_id, req=req, ngay=k.ngay, dong=[dict(
                hang_loai=k.hang_loai, hang_id=k.hang_id, so_luong=sl, sl_goc=tang, don_gia=gia,
                vi_tri=k.vi_tri if k.vi_tri is not None else k.vi_tri_hien, hsd=k.hsd,
                don_vi_goc=k.goc, kho_rong=k.kho_rong, kho_dai=k.kho_dai, ghi_chu=k.ghi_chu,
            )])
            self._ghi_so(v, req, user)
            for ln in v.lines:
                lot = self.db.get(StockLot, ln.lot_id)
                if lot is not None:
                    lot.ncc = k.ncc
        elif k.sl_goc < k.sl_hien - EPS:
            giam = k.sl_hien - k.sl_goc
            dong: list[dict] = []
            for lot in k.lots:
                if giam <= EPS:
                    break
                co = float(lot.sl_con_lai)
                if co <= EPS:
                    continue
                lay = min(co, giam)
                giam -= lay
                dong.append(dict(hang_loai=k.hang_loai, hang_id=k.hang_id, lot_id=lot.id, so_luong=lay,
                                 sl_goc=lay, don_gia=int(lot.don_gia_nhap or 0), don_vi_goc=k.goc,
                                 kho_rong=k.kho_rong, kho_dai=k.kho_dai))
            req = self._yeu_cau(user=user, kho_id=kho_id, loai="XUAT", k=k, dvt=k.goc,
                                sl=k.sl_hien - k.sl_goc, gia=None)
            v = self._phieu(user=user, kho_id=kho_id, req=req, ngay=k.ngay, dong=dong)
            self._ghi_so(v, req, user)
        if k.doi_vi_tri:
            for lot in self.lo_cua_dong(kho_id, k.khoa):
                if float(lot.sl_con_lai) > EPS:
                    lot.vi_tri = k.vi_tri

    def _dong_tang(self, k: KeHoach, tang: float) -> tuple[str, float, int]:
        """(đvt, số lượng, đơn giá) của dòng phiếu nhập phần tăng `tang` (đơn vị gốc). Ghi theo đơn vị
        mà giá là số tròn — giá lô = đơn giá × SL ÷ SL gốc mới ra đúng tiền: 100 kg × 21.000 =
        2.100.000, còn quy về 774 tờ × 2.713 đ là 2.099.862. Không ghi giá ⇒ theo đơn vị + giá của dòng
        nhập lô gần nhất. Số theo đơn vị đó không gọn 2 số lẻ (cột `so_luong`) thì mới ghi theo gốc."""
        ung_vien: list[tuple[str, float, float]] = []   # (đvt, số gốc / 1 đvt, đơn giá / đvt)
        if k.don_gia is not None:
            ung_vien.append((k.dvt, k.he_so, float(k.don_gia)))
        else:
            moi = max(k.lots, key=lambda x: (x.ngay_nhap, x.id))
            v = self.db.get(StockVoucher, moi.voucher_id) if moi.voucher_id else None
            ln = next((x for x in (v.lines if v is not None else []) if x.lot_id == moi.id), None)
            rl = self.db.get(StockRequestLine, ln.request_line_id) if ln is not None else None
            if rl is not None and float(ln.so_luong or 0) > 0 and ln.don_gia is not None:
                ung_vien.append((rl.dvt, float(ln.sl_goc) / float(ln.so_luong), float(ln.don_gia)))
            ung_vien.append((k.dvt, k.he_so, float(moi.don_gia_nhap or 0) * k.he_so))
        for dvt, he_so, gia in ung_vien:
            sl = tang / he_so if he_so else 0.0
            if round(sl, 2) > 0 and abs(sl - round(sl, 2)) < 1e-6:
                return dvt, round(sl, 2), round(gia)
        dvt, he_so, gia = ung_vien[0]
        return k.goc, tang, round(gia / he_so if he_so else gia)

    def _ghi_so(self, v: StockVoucher, req: StockRequest, user) -> None:
        from .stock_voucher_service import StockVoucherError

        self.db.refresh(v)
        try:
            self.v._apply_post(v, user)
        except StockVoucherError as e:
            raise KhoNhapTayError(str(e)) from None
        req.trang_thai = REQ_DONE
        self.db.flush()

    # ---------- một lượt (form / file) ----------

    def chay(self, kho_id: int, dong: list[DongVao], user, *, ghi: bool) -> KetQua:
        """Kiểm hết các dòng → còn lỗi thì dừng TRƯỚC khi chạm DB → không lỗi thì ghi cả lượt trong
        MỘT giao dịch. `ghi=False` = XEM TRƯỚC: chạy y hệt rồi rollback (như import khách hàng)."""
        self.kho(kho_id)
        kq = KetQua()
        vt = self.vi_tri_kho(kho_id)
        vi_tri_hop = set(vt) if vt else None
        ke_hoach: list[KeHoach] = []
        theo_khoa: dict[tuple, list[int]] = {}
        for d in dong:
            kq.tong_dong += 1
            k = self.doc_dong(kho_id, d, kq, vi_tri_hop=vi_tri_hop)
            if k is None:
                continue
            theo_khoa.setdefault(k.khoa, []).append(k.dong)
            ke_hoach.append(k)
        for khoa, cac_dong in theo_khoa.items():
            if len(cac_dong) > 1:
                for so in cac_dong:
                    kq.loi.append(LoiDong(so, NHAN["ma"], (
                        "Mặt hàng này (cùng mã, khổ, ĐVT) nằm ở nhiều dòng: "
                        + ", ".join(str(x) for x in cac_dong) + " — mỗi mặt hàng chỉ ghi một dòng.")))
        lam: list[KeHoach] = []
        for k in ke_hoach:
            if len(theo_khoa[k.khoa]) > 1:
                continue
            self.so_voi_kho(kho_id, k, kq)
            if not self.kiem_sau_so(kho_id, k, kq):
                continue
            if k.co_dong and not k.doi_so and not k.doi_vi_tri:
                kq.khong_doi += 1
                continue
            self.ghi_thay_doi(k, kq)
            if k.co_dong and not k.doi_so and k.don_gia is not None:
                kq.canh_bao.append(CanhBao(k.dong, "Số lượng không đổi nên đơn giá không được dùng — "
                                           "đơn giá chỉ áp cho phần tăng thêm."))
            if k.co_dong and k.sl_goc < k.sl_hien - EPS:
                kq.canh_bao.append(CanhBao(k.dong, (
                    f"{k.ma} giảm {_hien_so(k.sl_hien - k.sl_goc)} {k.goc_ten} — nếu lệnh nào đang giữ "
                    "chỗ hàng này, kế hoạch vật tư sẽ tính lại.")))
            lam.append(k)
        kq.loi.sort(key=lambda x: (x.dong, x.cot))
        kq.canh_bao.sort(key=lambda x: x.dong)
        if kq.loi:
            return kq

        so_dong = 0
        try:
            for k in lam:
                so_dong = k.dong
                if k.co_dong:
                    kq.cap_nhat += 1
                else:
                    kq.tao_moi += 1
                self.ap(kho_id, k, user)
            if lam:
                AuditLogRepository(self.db).create(
                    actor_user_id=getattr(user, "id", None), action="kho_hang_nhap_tay",
                    target=f"kho:{kho_id}",
                    detail=(f"Nhập tay tồn kho {self.kho(kho_id).ten}: {kq.tao_moi} thêm mới, "
                            f"{kq.cap_nhat} cập nhật — " + ", ".join(sorted({k.ma for k in lam}))),
                    commit=False,
                )
            if ghi:
                self.db.commit()
                kq.da_ghi = True
                if lam:
                    AuditLogRepository.bao_co_dong_moi()
            else:
                self.db.rollback()
        except KhoNhapTayError as e:
            self.db.rollback()
            kq.loi.append(LoiDong(so_dong, "", str(e)))
            kq.tao_moi = kq.cap_nhat = 0
        except Exception:
            self.db.rollback()
            raise
        return kq

    # ---------- xoá ----------

    def xoa(self, kho_id: int, khoas: list[tuple], user) -> int:
        """Xoá hẳn các dòng (mã + khổ + đơn vị gốc) khỏi kho. Trả số lô đã xoá."""
        kho = self.kho(kho_id)
        lot_ids: list[int] = []
        ten: list[str] = []
        for khoa in khoas:
            lots = self.lo_cua_dong(kho_id, khoa)
            if not lots:
                raise KhoNhapTayError("Dòng cần xoá không còn trong kho — tải lại màn hình.")
            lot_ids.extend(x.id for x in lots)
            m = self.hang.map_theo_cap([(khoa[0], khoa[1])]).get((khoa[0], khoa[1]))
            ten.append(getattr(m, "ma", "?"))
        if not lot_ids:
            return 0
        dong_phieu = list(self.db.execute(
            select(StockVoucherLine, StockVoucher)
            .join(StockVoucher, StockVoucher.id == StockVoucherLine.voucher_id)
            .where(StockVoucherLine.lot_id.in_(lot_ids))).all())
        # Lô nhập (phiếu tạo ra lô) cũng là lịch sử của dòng — dòng phiếu nhập không trỏ `lot_id`
        # trước khi ghi sổ, nhưng sau ghi sổ thì có (`ln.lot_id = lot.id`), nên câu trên đã gom đủ.
        for ln, v in dong_phieu:
            if v.trang_thai == "posted" and v.ngay:
                chan = self._chan_ky(kho_id, v.ngay)
                if chan:
                    raise KhoNhapTayError(
                        f"Không xoá được: phiếu {v.ma} của mặt hàng này có ngày nằm trong kỳ đã chốt. {chan}")
        phieu_ids = {v.id for _ln, v in dong_phieu}
        self.db.execute(delete(StockVoucherLine).where(StockVoucherLine.lot_id.in_(lot_ids)))
        for model in (StockLot, StockVoucherLine, StockRequestLine):
            self.db.execute(update(model).where(model.lo_goc_id.in_(lot_ids)).values(lo_goc_id=None))
        self.db.execute(delete(StockLot).where(StockLot.id.in_(lot_ids)))
        for khoa in khoas:
            loai, hid, kr, kd, goc = khoa
            self.db.execute(delete(StockThreshold).where(
                StockThreshold.kho_id == kho_id, StockThreshold.hang_loai == loai,
                StockThreshold.hang_id == hid, StockThreshold.kho_rong == kr,
                StockThreshold.kho_dai == kd, StockThreshold.don_vi_goc == goc))
        # Phiếu nhập tay chỉ còn rỗng ⇒ xoá luôn phiếu + yêu cầu tự sinh của nó.
        for vid in phieu_ids:
            v = self.db.get(StockVoucher, vid)
            if v is None:
                continue
            con = self.db.execute(select(StockVoucherLine.id).where(StockVoucherLine.voucher_id == vid)
                                  .limit(1)).first()
            req = self.db.get(StockRequest, v.request_id) if v.request_id else None
            if con is None and req is not None and req.nhap_tay:
                self.db.delete(v)
                self.db.flush()
                con_phieu = self.db.execute(select(StockVoucher.id).where(StockVoucher.request_id == req.id)
                                            .limit(1)).first()
                if con_phieu is None:
                    self.db.delete(req)
        AuditLogRepository(self.db).create(
            actor_user_id=getattr(user, "id", None), action="kho_hang_xoa", target=f"kho:{kho_id}",
            detail=f"Xoá khỏi {kho.ten}: " + ", ".join(ten), commit=False,
        )
        self.db.commit()
        AuditLogRepository.bao_co_dong_moi()
        return len(lot_ids)

    # ---------- dòng tồn (export) ----------

    def dong_ton(self, kho_id: int) -> list[dict]:
        """Mỗi dòng = mã + khổ + đơn vị gốc còn hàng. Đơn giá = BQ các lô còn tồn; vị trí, NCC, ngày
        nhập theo lô nhập gần nhất; hạn = hạn sớm nhất của lô còn tồn."""
        lots = list(self.db.execute(
            select(StockLot).where(StockLot.kho_id == kho_id, StockLot.sl_con_lai > 0)
            .where(StockLot.hang_loai.in_(("giay", "vat_tu")))
        ).scalars())
        gom: dict[tuple, list[StockLot]] = {}
        for x in lots:
            gom.setdefault((x.hang_loai, x.hang_id, x.kho_rong, x.kho_dai, x.don_vi_goc or ""), []).append(x)
        hang = self.hang.map_theo_cap({(k[0], k[1]) for k in gom})
        ra = []
        for khoa, ds in gom.items():
            m = hang.get((khoa[0], khoa[1]))
            if m is None or getattr(m, "la_thanh_pham", False):
                continue
            sl = sum(float(x.sl_con_lai) for x in ds)
            gt = sum(float(x.sl_con_lai) * float(x.don_gia_nhap or 0) for x in ds)
            moi = max(ds, key=lambda x: (x.ngay_nhap, x.id))
            hsd = [x.hsd for x in ds if x.hsd]
            goc = khoa[4] or getattr(m, "don_vi_gia", "") or ""
            ra.append(dict(
                ma=m.ma, ten=m.ten, dinh_luong=getattr(m, "gsm", None) if khoa[0] == "giay" else None,
                kho_rong=khoa[2] or None, kho_dai=khoa[3] or None, dvt=self.ten_dv(goc),
                so_luong=sl, don_gia=round(gt / sl, 3) if sl else 0, vi_tri=moi.vi_tri,
                hsd=min(hsd) if hsd else None, ngay_nhap=moi.ngay_nhap, ncc=moi.ncc, ghi_chu=None,
            ))
        ra.sort(key=lambda r: (r["ma"].upper(), r["kho_rong"] or 0, r["kho_dai"] or 0, r["dvt"]))
        return ra

    # ---------- Excel ----------

    def _dvt_chon(self) -> list[str]:
        return sorted({d.ten for d in self.hang.don_vi.all_active() if d.ten})

    def tao_file(self, kho_id: int, rows: list[dict] | None) -> bytes:
        """File mẫu (rows=None, 200 dòng kẻ sẵn) hoặc file xuất. Cùng cột, cùng dấu `_meta`."""
        kho = self.kho(kho_id)
        font, vien = "Arial", Side(style="thin", color="BFBFBF")
        ke = Border(left=vien, right=vien, top=vien, bottom=vien)
        nen_dau = PatternFill("solid", fgColor="E7E6E6")
        chu, chu_dau = Font(name=font, size=10), Font(name=font, size=10, bold=True)
        idx = {c.khoa: i for i, c in enumerate(COT, start=1)}
        L = {k: get_column_letter(i) for k, i in idx.items()}

        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        for i, c in enumerate(COT, start=1):
            o = ws.cell(row=1, column=i, value=c.nhan + (" (*)" if c.bat_buoc else ""))
            o.font, o.fill, o.border = chu_dau, nen_dau, ke
            o.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = c.rong + 4
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = "A2"

        so = len(rows) if rows is not None else 200
        for r in range(2, so + 2):
            for i, c in enumerate(COT, start=1):
                o = ws.cell(row=r, column=i)
                o.font, o.border, o.number_format = chu, ke, c.dinh_dang
                o.alignment = Alignment(horizontal=c.can, vertical="center")
            g, h = L["so_luong"], L["don_gia"]
            ws.cell(row=r, column=idx["thanh_tien"],
                    value=f'=IF(OR({g}{r}="",{h}{r}=""),"",ROUND({g}{r}*{h}{r},0))')
        for r, row in enumerate(rows or [], start=2):
            for c in COT:
                if c.khoa == "thanh_tien" or c.khoa not in row:
                    continue
                ws.cell(row=r, column=idx[c.khoa], value=row[c.khoa])
            for k in ("so_luong", "don_gia"):
                o = ws.cell(row=r, column=idx[k])
                if isinstance(o.value, (int, float)):
                    o.number_format = "#,##0" if float(o.value).is_integer() else "#,##0.###"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COT))}{max(ws.max_row, 2)}"

        ds = wb.create_sheet(SHEET_DS)
        dvt = self._dvt_chon()
        vt = self.vi_tri_kho(kho_id)
        for i, v in enumerate(dvt, start=1):
            ds.cell(row=i, column=1, value=v)
        for i, v in enumerate(vt, start=1):
            ds.cell(row=i, column=2, value=v)
        ds.sheet_state = "hidden"
        han = max(so + 1, 1000)
        if dvt:
            dv = DataValidation(type="list", formula1=f"={SHEET_DS}!$A$1:$A${len(dvt)}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{L['dvt']}2:{L['dvt']}{han}")
        if vt:
            dv = DataValidation(type="list", formula1=f"={SHEET_DS}!$B$1:$B${len(vt)}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{L['vi_tri']}2:{L['vi_tri']}{han}")
        for k in ("so_luong", "don_gia", "kho_rong", "kho_dai"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            dv.error = "Nhập số lớn hơn hoặc bằng 0."
            ws.add_data_validation(dv)
            dv.add(f"{L[k]}2:{L[k]}{han}")
        for k in ("hsd", "ngay_nhap"):
            dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
            dv.error = "Nhập ngày theo dạng dd/mm/yyyy."
            ws.add_data_validation(dv)
            dv.add(f"{L[k]}2:{L[k]}{han}")

        hd = wb.create_sheet(SHEET_HD)
        for c, w in (("A", 6), ("B", 16), ("C", 10), ("D", 80)):
            hd.column_dimensions[c].width = w
        hd["A1"] = f"HƯỚNG DẪN NHẬP TỒN KHO — {kho.ten.upper()}"
        hd["A1"].font = Font(name=font, size=12, bold=True)
        hd["A2"] = ("Cột có dấu (*) là bắt buộc. Tên hàng, Định lượng, Thành tiền do hệ thống tự điền, ghi vào "
                    "sẽ không được dùng.")
        hd["A2"].font = chu
        for i, t in enumerate(("STT", "Cột", "Bắt buộc", "Cách ghi"), start=1):
            o = hd.cell(row=4, column=i, value=t)
            o.font, o.fill, o.border = chu_dau, nen_dau, ke
            o.alignment = Alignment(horizontal="center", vertical="center")
        r = 5
        for stt, c in enumerate(COT, start=1):
            for i, v in enumerate((stt, c.nhan, "x" if c.bat_buoc else "", c.cach_ghi), start=1):
                o = hd.cell(row=r, column=i, value=v)
                o.font, o.border = chu, ke
                o.alignment = Alignment(horizontal="center" if i in (1, 3) else "left", vertical="center")
            r += 1
        r += 1
        hd.cell(row=r, column=1, value="Lưu ý").font = chu_dau
        r += 1
        for stt, luat in enumerate(LUU_Y, start=1):
            hd.cell(row=r, column=1, value=f"{stt}.").font = chu
            hd.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
            hd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            o = hd.cell(row=r, column=2, value=luat)
            o.font = chu
            o.alignment = Alignment(wrap_text=True, vertical="top")
            hd.row_dimensions[r].height = 27 if len(luat) > 135 else 14
            r += 1

        meta = wb.create_sheet(SHEET_META)
        for d in (("khoa", "gia_tri"), ("loai", LOAI_META), ("kho_id", kho.id), ("kho_ma", kho.ma),
                  ("phien_ban", PHIEN_BAN)):
            meta.append(d)
        meta.sheet_state = "hidden"
        wb.active = 0
        ra = BytesIO()
        wb.save(ra)
        return ra.getvalue()

    def doc_file(self, kho_id: int, du_lieu: bytes) -> list[DongVao]:
        """File → các dòng. Lỗi cả file ⇒ `ExcelSaiMan` (từ chối ngay, không xem trước)."""
        if not du_lieu:
            raise ExcelSaiMan("Không đọc được file — phải là .xlsx đúng mẫu tải từ hệ thống.")
        try:
            wb = load_workbook(BytesIO(du_lieu), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 — mọi kiểu hỏng của file đều cùng một câu trả lời
            raise ExcelSaiMan("Không đọc được file — phải là .xlsx đúng mẫu tải từ hệ thống.") from None
        if SHEET_META not in wb.sheetnames:
            raise ExcelSaiMan('File không phải file tồn kho (thiếu dấu "_meta"). Dùng file tải từ nút '
                              '"Tải file mẫu" hoặc "Export" của màn Tồn kho.')
        gia: dict[str, Any] = {}
        for hang in wb[SHEET_META].iter_rows(min_row=2, values_only=True):
            if hang and hang[0]:
                gia[str(hang[0]).strip()] = hang[1]
        if str(gia.get("loai") or "").strip() != LOAI_META:
            raise ExcelSaiMan(f'File này là mẫu của màn khác ("{gia.get("loai")}"), không phải Tồn kho.')
        try:
            kho_file = int(gia.get("kho_id"))
        except (TypeError, ValueError):
            kho_file = None
        if kho_file != kho_id:
            raise ExcelSaiMan(f"File này của kho {gia.get('kho_ma') or '?'} — chỉ nhập được vào đúng kho đã "
                              f"tải file về. Tải file mẫu của kho {self.kho(kho_id).ten}.")
        if SHEET not in wb.sheetnames:
            raise ExcelSaiMan(f'File thiếu sheet "{SHEET}".')
        hang = list(wb[SHEET].iter_rows(values_only=True))
        if not hang:
            raise ExcelSaiMan(f'Sheet "{SHEET}" trống — không có cả dòng tiêu đề.')
        tieu_de = [_nhan_goc(_chu(o) or "") for o in hang[0]]
        vi_tri = {c.khoa: tieu_de.index(c.nhan.lower()) for c in COT if c.nhan.lower() in tieu_de}
        thieu = [c.nhan for c in COT if c.bat_buoc and c.khoa not in vi_tri]
        if thieu:
            raise ExcelSaiMan("File thiếu cột bắt buộc: " + ", ".join(thieu) + ". Tải lại file mẫu.")
        ra: list[DongVao] = []
        for so, dong in enumerate(hang[1:], start=2):
            o = {k: (dong[i] if i < len(dong) else None) for k, i in vi_tri.items()}
            # Dòng chỉ có công thức Thành tiền (mẫu kẻ sẵn) hoặc trắng hoàn toàn ⇒ bỏ qua, không là lỗi.
            if all(_chu(v) is None for k, v in o.items() if k not in ("thanh_tien", "ten", "dinh_luong")):
                continue
            ra.append(DongVao(dong=so, **{k: o.get(k) for k in (
                "ma", "dvt", "so_luong", "don_gia", "kho_rong", "kho_dai", "vi_tri", "hsd", "ngay_nhap",
                "ncc", "ghi_chu")}))
        return ra

    # ---------- xem theo ĐVT (spec 7b) ----------

    def don_vi_xem(self, lots) -> dict[tuple, list[dict]]:
        """`{(loai, id, rộng, dài, gốc): [{ma, ten, he_so}]}` — các đơn vị QUY ĐƯỢC từ đơn vị gốc của
        dòng, `he_so` = số đơn vị gốc trong 1 đơn vị đó. Đơn vị gốc đứng đầu. Giấy tờ khổ hai cạnh có
        định lượng thêm kg (quy qua gsm). Không quy được thì không có mặt."""
        khoa = {(x.hang_loai, x.hang_id, int(x.kho_rong or 0), int(x.kho_dai or 0), (x.don_vi_goc or "").lower())
                for x in lots if x.hang_loai in ("giay", "vat_tu") and x.don_vi_goc}
        if not khoa:
            return {}
        nhom = self.hang.don_vi_nhieu_theo_goc({(k[0], k[1], k[4]) for k in khoa})
        hang = self.hang.map_theo_cap({(k[0], k[1]) for k in khoa})
        ra: dict[tuple, list[dict]] = {}
        for k in khoa:
            g = nhom.get((k[0], k[1], k[4]))
            ds: list[dict] = [{"ma": k[4], "ten": self.ten_dv(k[4]), "he_so": 1.0}]
            if g:
                for d in g.get("ds") or []:
                    if (d.get("goc") or g.get("don_vi_goc")) != k[4] or d["ma"] == k[4]:
                        continue
                    hs = float(d.get("he_so_ve_goc") or 0)
                    if hs > 0:
                        ds.append({"ma": d["ma"], "ten": d.get("ten") or d["ma"], "he_so": hs})
            m = hang.get((k[0], k[1]))
            gsm = getattr(m, "gsm", None) if k[0] == "giay" else None
            if k[0] == "giay" and k[4] == DON_VI_TO_KHO and la_kho_to(k[2], k[3]) and gsm:
                kg_1to = kg_tu_to(1, k[2], k[3], gsm)
                if kg_1to > 0 and not any(x["ma"] == "kg" for x in ds):
                    ds.append({"ma": "kg", "ten": self.ten_dv("kg") or "kg", "he_so": 1 / kg_1to})
            ra[k] = ds
        return ra
